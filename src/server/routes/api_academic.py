import io
import logging
from typing import Optional, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fastapi import APIRouter, Depends, HTTPException, Request, status, Query, UploadFile, File
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.orm import Session
from sqlalchemy import or_

from src.database.models import (
    Tenant,
    User,
    Department,
    AcademicYear,
    ClassModel,
    Division,
    TeacherClassAssignment,
    Student,
    StudentBatchUpload,
    AuditLog,
)
from src.database.session import get_db
from src.server.tenant_middleware import get_current_tenant
from src.server.rbac_middleware import require_roles, get_current_user, check_tenant_operational_access
from src.utils.auth_utils import hash_password
from src.utils.timezone import get_ist_now
from src.core.face_engine import face_engine

logger = logging.getLogger("api_academic")
router = APIRouter(
    prefix="/api/v1/academic",
    tags=["Academic Hierarchy & Progression"],
    dependencies=[Depends(require_roles(["TENANT_ADMIN"]))],
)


# --- Request & Response Schemas ---
class CreateDepartmentRequest(BaseModel):
    name: str
    code: Optional[str] = None
    description: Optional[str] = None


class UpdateDepartmentRequest(BaseModel):
    name: Optional[str] = None
    code: Optional[str] = None
    description: Optional[str] = None


class CreateAcademicYearRequest(BaseModel):
    name: str
    is_current: bool = True
    start_date: Optional[str] = None
    end_date: Optional[str] = None


class CreateClassRequest(BaseModel):
    name: str
    department_id: Optional[int] = None
    department: Optional[str] = "Computer Science"
    code: Optional[str] = None


class UpdateClassRequest(BaseModel):
    name: Optional[str] = None
    department_id: Optional[int] = None
    department: Optional[str] = None
    code: Optional[str] = None


class CreateDivisionRequest(BaseModel):
    class_id: int
    name: str


class UpdateDivisionRequest(BaseModel):
    name: Optional[str] = None
    class_id: Optional[int] = None


class CreateTeacherUserRequest(BaseModel):
    username: str
    email: str
    full_name: str
    password: str
    phone_number: Optional[str] = None


class AssignTeacherRequest(BaseModel):
    teacher_id: int
    class_id: int
    division_id: Optional[int] = None
    academic_year_id: Optional[int] = None
    subject: str = "General"


class StudentPromotionRequest(BaseModel):
    student_ids: List[int]
    target_class_id: int
    target_division_id: Optional[int] = None
    target_academic_year_id: Optional[int] = None


class PromotionRollbackRequest(BaseModel):
    student_ids: Optional[List[int]] = None
    class_id: Optional[int] = None  # Rollback whole class if specified


class StudentTransferRequest(BaseModel):
    student_id: int
    target_department_id: int
    target_class_id: int
    target_division_id: Optional[int] = None
    target_academic_year_id: Optional[int] = None


class StudentTransferRollbackRequest(BaseModel):
    student_id: int


class BulkStudentImportItem(BaseModel):
    name: str
    roll_number: str
    department_id: int
    class_id: int
    division_id: Optional[int] = None
    gender: Optional[str] = "Other"
    user_role: Optional[str] = "student"
    email: Optional[str] = None
    phone_number: Optional[str] = None


class BulkStudentImportRequest(BaseModel):
    rows: List[BulkStudentImportItem]
    filename: Optional[str] = "bulk_upload.xlsx"
    academic_year_id: Optional[int] = None


# ==========================================================
# 1. DEPARTMENTS CRUD & REFERENTIAL INTEGRITY
# ==========================================================
@router.get("/departments")
def list_departments(
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """Lists all academic departments for current tenant."""
    departments = (
        db.query(Department)
        .filter(Department.tenant_id == tenant.id)
        .order_by(Department.name.asc())
        .all()
    )
    return {"status": "success", "departments": [d.to_dict() for d in departments]}


@router.post("/departments")
def create_department(
    payload: CreateDepartmentRequest,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """Creates a new institutional Department scoped to the active tenant."""
    check_tenant_operational_access(tenant)
    name_clean = payload.name.strip()
    if not name_clean:
        raise HTTPException(status_code=400, detail="Department name cannot be empty.")

    existing = db.query(Department).filter(
        Department.tenant_id == tenant.id,
        Department.name == name_clean,
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Department '{name_clean}' already exists in this institution.")

    code_clean = payload.code.strip().upper() if payload.code else name_clean[:6].upper()
    dept = Department(
        tenant_id=tenant.id,
        name=name_clean,
        code=code_clean,
        description=payload.description.strip() if payload.description else "",
    )
    db.add(dept)
    db.commit()
    db.refresh(dept)
    return {"status": "success", "message": f"Department '{name_clean}' created.", "department": dept.to_dict()}


@router.put("/departments/{department_id}")
def update_department(
    department_id: int,
    payload: UpdateDepartmentRequest,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """Updates an existing department's name, code, and description."""
    check_tenant_operational_access(tenant)
    dept = db.query(Department).filter(
        Department.tenant_id == tenant.id,
        Department.id == department_id,
    ).first()
    if not dept:
        raise HTTPException(status_code=404, detail="Department not found.")

    if payload.name:
        name_clean = payload.name.strip()
        if name_clean != dept.name:
            dup = db.query(Department).filter(
                Department.tenant_id == tenant.id,
                Department.name == name_clean,
                Department.id != department_id,
            ).first()
            if dup:
                raise HTTPException(status_code=400, detail=f"Another department named '{name_clean}' already exists.")
            
            # Update matching legacy text references on classes and students
            db.query(ClassModel).filter(
                ClassModel.tenant_id == tenant.id,
                ClassModel.department_id == department_id,
            ).update({"department": name_clean})

            db.query(Student).filter(
                Student.tenant_id == tenant.id,
                Student.department_id == department_id,
            ).update({"department": name_clean})

            dept.name = name_clean

    if payload.code is not None:
        dept.code = payload.code.strip().upper()
    if payload.description is not None:
        dept.description = payload.description.strip()

    db.commit()
    db.refresh(dept)
    return {"status": "success", "message": f"Department '{dept.name}' updated.", "department": dept.to_dict()}


@router.delete("/departments/{department_id}")
def delete_department(
    department_id: int,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """
    Deletes an empty department.
    Enforces strict referential integrity (RESTRICT): rejects deletion if active classes or students are linked.
    """
    check_tenant_operational_access(tenant)
    dept = db.query(Department).filter(
        Department.tenant_id == tenant.id,
        Department.id == department_id,
    ).first()
    if not dept:
        raise HTTPException(status_code=404, detail="Department not found.")

    classes_count = db.query(ClassModel).filter(
        ClassModel.tenant_id == tenant.id,
        or_(ClassModel.department_id == department_id, ClassModel.department == dept.name),
    ).count()

    students_count = db.query(Student).filter(
        Student.tenant_id == tenant.id,
        or_(Student.department_id == department_id, Student.department == dept.name),
    ).count()

    if classes_count > 0 or students_count > 0:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Cannot delete Department '{dept.name}': It contains {classes_count} class(es) "
                f"and {students_count} student(s). Please reassign or delete these records first."
            ),
        )

    dept_name = dept.name
    db.delete(dept)
    db.commit()
    return {"status": "success", "message": f"Department '{dept_name}' deleted successfully."}


# ==========================================================
# 2. ACADEMIC YEARS
# ==========================================================
@router.get("/years")
def list_academic_years(
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    years = db.query(AcademicYear).filter(AcademicYear.tenant_id == tenant.id).order_by(AcademicYear.id.desc()).all()
    return {"status": "success", "years": [y.to_dict() for y in years]}


@router.post("/years")
def create_academic_year(
    payload: CreateAcademicYearRequest,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    check_tenant_operational_access(tenant)
    name_clean = payload.name.strip()
    if payload.is_current:
        # De-activate previous current years
        db.query(AcademicYear).filter(AcademicYear.tenant_id == tenant.id).update({"is_current": False})

    existing = db.query(AcademicYear).filter(AcademicYear.tenant_id == tenant.id, AcademicYear.name == name_clean).first()
    if existing:
        existing.is_current = payload.is_current
        db.commit()
        return {"status": "success", "message": f"Academic Year '{name_clean}' updated.", "year": existing.to_dict()}

    year = AcademicYear(
        tenant_id=tenant.id,
        name=name_clean,
        is_current=payload.is_current,
        start_date=payload.start_date,
        end_date=payload.end_date,
    )
    db.add(year)
    db.commit()
    db.refresh(year)
    return {"status": "success", "message": f"Academic Year '{name_clean}' created.", "year": year.to_dict()}


# ==========================================================
# 3. CLASSES & REFERENTIAL INTEGRITY
# ==========================================================
@router.get("/classes")
def list_classes(
    department_id: Optional[int] = None,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    query = db.query(ClassModel).filter(ClassModel.tenant_id == tenant.id)
    if department_id:
        query = query.filter(ClassModel.department_id == department_id)
    classes = query.order_by(ClassModel.id.asc()).all()
    return {"status": "success", "classes": [c.to_dict() for c in classes]}


@router.post("/classes")
def create_class(
    payload: CreateClassRequest,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    check_tenant_operational_access(tenant)
    name_clean = payload.name.strip()
    existing = db.query(ClassModel).filter(ClassModel.tenant_id == tenant.id, ClassModel.name == name_clean).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Class '{name_clean}' already exists in this institution.")

    # Resolve department
    dept_id = payload.department_id
    dept_name = (payload.department or "Computer Science").strip()

    if dept_id:
        dept_obj = db.query(Department).filter(Department.tenant_id == tenant.id, Department.id == dept_id).first()
        if not dept_obj:
            raise HTTPException(status_code=404, detail="Selected department not found.")
        dept_name = dept_obj.name
    else:
        dept_obj = db.query(Department).filter(Department.tenant_id == tenant.id, Department.name == dept_name).first()
        if dept_obj:
            dept_id = dept_obj.id

    class_obj = ClassModel(
        tenant_id=tenant.id,
        department_id=dept_id,
        department=dept_name,
        name=name_clean,
        code=payload.code.strip() if payload.code else name_clean[:10].upper(),
    )
    db.add(class_obj)
    db.commit()
    db.refresh(class_obj)
    return {"status": "success", "message": f"Class '{name_clean}' created.", "class": class_obj.to_dict()}


@router.put("/classes/{class_id}")
def update_class(
    class_id: int,
    payload: UpdateClassRequest,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """Updates class name, code, or department assignment."""
    check_tenant_operational_access(tenant)
    class_obj = db.query(ClassModel).filter(
        ClassModel.tenant_id == tenant.id,
        ClassModel.id == class_id,
    ).first()
    if not class_obj:
        raise HTTPException(status_code=404, detail="Class not found.")

    if payload.name:
        name_clean = payload.name.strip()
        if name_clean != class_obj.name:
            dup = db.query(ClassModel).filter(
                ClassModel.tenant_id == tenant.id,
                ClassModel.name == name_clean,
                ClassModel.id != class_id,
            ).first()
            if dup:
                raise HTTPException(status_code=400, detail=f"Another class named '{name_clean}' already exists.")
            class_obj.name = name_clean

    if payload.department_id:
        dept_obj = db.query(Department).filter(
            Department.tenant_id == tenant.id,
            Department.id == payload.department_id,
        ).first()
        if not dept_obj:
            raise HTTPException(status_code=404, detail="Selected department not found.")
        class_obj.department_id = dept_obj.id
        class_obj.department = dept_obj.name
    elif payload.department:
        dept_clean = payload.department.strip()
        dept_obj = db.query(Department).filter(
            Department.tenant_id == tenant.id,
            Department.name == dept_clean,
        ).first()
        if dept_obj:
            class_obj.department_id = dept_obj.id
            class_obj.department = dept_obj.name
        else:
            class_obj.department = dept_clean

    if payload.code is not None:
        class_obj.code = payload.code.strip()

    db.commit()
    db.refresh(class_obj)
    return {"status": "success", "message": f"Class '{class_obj.name}' updated.", "class": class_obj.to_dict()}


@router.delete("/classes/{class_id}")
def delete_class(
    class_id: int,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """
    Deletes a class created by mistake.
    Enforces strict referential integrity (RESTRICT): rejects deletion if active students are assigned.
    """
    check_tenant_operational_access(tenant)
    class_obj = db.query(ClassModel).filter(
        ClassModel.tenant_id == tenant.id,
        ClassModel.id == class_id,
    ).first()
    if not class_obj:
        raise HTTPException(status_code=404, detail="Class not found.")

    students_count = db.query(Student).filter(
        Student.tenant_id == tenant.id,
        Student.class_id == class_id,
    ).count()

    if students_count > 0:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Cannot delete Class '{class_obj.name}': It contains {students_count} enrolled student(s). "
                f"Please promote or reassign these students before deleting the class."
            ),
        )

    class_name = class_obj.name
    db.delete(class_obj)
    db.commit()
    return {"status": "success", "message": f"Class '{class_name}' and associated divisions deleted."}


# ==========================================================
# 4. DIVISIONS & REFERENTIAL INTEGRITY
# ==========================================================
@router.get("/divisions")
def list_divisions(
    class_id: Optional[int] = None,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    query = db.query(Division).filter(Division.tenant_id == tenant.id)
    if class_id:
        query = query.filter(Division.class_id == class_id)
    divisions = query.order_by(Division.id.asc()).all()
    return {"status": "success", "divisions": [d.to_dict() for d in divisions]}


@router.post("/divisions")
def create_division(
    payload: CreateDivisionRequest,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    check_tenant_operational_access(tenant)
    name_clean = payload.name.strip()
    class_obj = db.query(ClassModel).filter(ClassModel.tenant_id == tenant.id, ClassModel.id == payload.class_id).first()
    if not class_obj:
        raise HTTPException(status_code=404, detail="Target class not found.")

    existing = db.query(Division).filter(
        Division.tenant_id == tenant.id,
        Division.class_id == payload.class_id,
        Division.name == name_clean,
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Division '{name_clean}' already exists for {class_obj.name}.")

    div = Division(
        tenant_id=tenant.id,
        class_id=payload.class_id,
        name=name_clean,
    )
    db.add(div)
    db.commit()
    db.refresh(div)
    return {"status": "success", "message": f"Division '{name_clean}' created.", "division": div.to_dict()}


@router.put("/divisions/{division_id}")
def update_division(
    division_id: int,
    payload: UpdateDivisionRequest,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """Updates division name or reassigns parent class."""
    check_tenant_operational_access(tenant)
    div = db.query(Division).filter(
        Division.tenant_id == tenant.id,
        Division.id == division_id,
    ).first()
    if not div:
        raise HTTPException(status_code=404, detail="Division not found.")

    if payload.name:
        name_clean = payload.name.strip()
        target_class_id = payload.class_id or div.class_id
        dup = db.query(Division).filter(
            Division.tenant_id == tenant.id,
            Division.class_id == target_class_id,
            Division.name == name_clean,
            Division.id != division_id,
        ).first()
        if dup:
            raise HTTPException(status_code=400, detail=f"Division '{name_clean}' already exists for this class.")
        div.name = name_clean

    if payload.class_id and payload.class_id != div.class_id:
        target_class = db.query(ClassModel).filter(
            ClassModel.tenant_id == tenant.id,
            ClassModel.id == payload.class_id,
        ).first()
        if not target_class:
            raise HTTPException(status_code=404, detail="Target parent class not found.")
        div.class_id = payload.class_id

    db.commit()
    db.refresh(div)
    return {"status": "success", "message": f"Division '{div.name}' updated.", "division": div.to_dict()}


@router.delete("/divisions/{division_id}")
def delete_division(
    division_id: int,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """
    Deletes a division.
    Enforces strict referential integrity (RESTRICT): rejects deletion if students are assigned to this division.
    """
    check_tenant_operational_access(tenant)
    div = db.query(Division).filter(
        Division.tenant_id == tenant.id,
        Division.id == division_id,
    ).first()
    if not div:
        raise HTTPException(status_code=404, detail="Division not found.")

    students_count = db.query(Student).filter(
        Student.tenant_id == tenant.id,
        Student.division_id == division_id,
    ).count()

    if students_count > 0:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Cannot delete Division '{div.name}': It contains {students_count} enrolled student(s). "
                f"Please reassign students to another division or parent class first."
            ),
        )

    div_name = div.name
    db.delete(div)
    db.commit()
    return {"status": "success", "message": f"Division '{div_name}' deleted successfully."}


# ==========================================================
# 5. TEACHERS & FACULTY ASSIGNMENTS
# ==========================================================
@router.get("/teachers")
def list_teachers(
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    teachers = db.query(User).filter(User.tenant_id == tenant.id, User.role == "TEACHER").all()
    result = []
    for t in teachers:
        t_dict = t.to_dict()
        t_dict["assignments"] = [a.to_dict() for a in (t.teacher_assignments or [])]
        result.append(t_dict)
    return {"status": "success", "teachers": result}


@router.post("/teachers")
def create_teacher(
    payload: CreateTeacherUserRequest,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    check_tenant_operational_access(tenant)
    username_clean = payload.username.strip().lower()
    existing = db.query(User).filter(User.tenant_id == tenant.id, User.username == username_clean).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Teacher username '{username_clean}' is already taken.")

    teacher = User(
        tenant_id=tenant.id,
        username=username_clean,
        email=payload.email.strip(),
        full_name=payload.full_name.strip(),
        password_hash=hash_password(payload.password),
        phone_number=payload.phone_number,
        role="TEACHER",
        is_active=True,
    )
    db.add(teacher)
    db.commit()
    db.refresh(teacher)
    return {"status": "success", "message": f"Teacher account '{teacher.full_name}' created.", "teacher": teacher.to_dict()}


@router.get("/teacher-assignments")
def list_teacher_assignments(
    teacher_id: Optional[int] = None,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    query = db.query(TeacherClassAssignment).filter(TeacherClassAssignment.tenant_id == tenant.id)
    if teacher_id:
        query = query.filter(TeacherClassAssignment.teacher_id == teacher_id)
    assignments = query.all()
    return {"status": "success", "assignments": [a.to_dict() for a in assignments]}


@router.post("/teacher-assignments")
def assign_teacher_to_class(
    payload: AssignTeacherRequest,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    check_tenant_operational_access(tenant)
    teacher = db.query(User).filter(User.tenant_id == tenant.id, User.id == payload.teacher_id, User.role == "TEACHER").first()
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher user not found.")

    class_obj = db.query(ClassModel).filter(ClassModel.tenant_id == tenant.id, ClassModel.id == payload.class_id).first()
    if not class_obj:
        raise HTTPException(status_code=404, detail="Target class not found.")

    assignment = TeacherClassAssignment(
        tenant_id=tenant.id,
        teacher_id=payload.teacher_id,
        class_id=payload.class_id,
        division_id=payload.division_id,
        academic_year_id=payload.academic_year_id,
        subject=payload.subject.strip(),
    )
    db.add(assignment)
    db.commit()
    db.refresh(assignment)
    return {"status": "success", "message": f"Assigned {teacher.full_name} to {class_obj.name}.", "assignment": assignment.to_dict()}


@router.delete("/teacher-assignments/{assignment_id}")
def remove_teacher_assignment(
    assignment_id: int,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    check_tenant_operational_access(tenant)
    assignment = db.query(TeacherClassAssignment).filter(
        TeacherClassAssignment.tenant_id == tenant.id,
        TeacherClassAssignment.id == assignment_id,
    ).first()
    if not assignment:
        raise HTTPException(status_code=404, detail="Assignment not found.")

    db.delete(assignment)
    db.commit()
    return {"status": "success", "message": "Teacher assignment removed."}


# ==========================================================
# 6. FILTERED STUDENT DIRECTORY & PROMOTION COHORTS
# ==========================================================
@router.get("/students")
def list_cohort_students(
    department_id: Optional[int] = None,
    class_id: Optional[int] = None,
    division_id: Optional[int] = None,
    academic_year_id: Optional[int] = None,
    role: Optional[str] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """
    Returns filtered student roster for academic progression, promotion engine, and directory.
    Eliminates infinite loading by supporting direct database query filters.
    """
    query = db.query(Student).filter(Student.tenant_id == tenant.id)

    if department_id:
        query = query.filter(Student.department_id == department_id)
    if class_id:
        query = query.filter(Student.class_id == class_id)
    if division_id:
        query = query.filter(Student.division_id == division_id)
    if academic_year_id:
        query = query.filter(Student.academic_year_id == academic_year_id)
    if role:
        query = query.filter(Student.user_role == role.strip().lower())
    if search:
        search_term = f"%{search.strip()}%"
        query = query.filter(
            or_(
                Student.name.ilike(search_term),
                Student.roll_number.ilike(search_term),
                Student.email.ilike(search_term),
            )
        )

    students = query.order_by(Student.roll_number.asc()).all()
    return {
        "status": "success",
        "students": [s.to_dict() for s in students],
        "count": len(students),
    }


# ==========================================================
# 7. STUDENT PROMOTION & ROLLBACK ENGINE
# ==========================================================
@router.post("/students/promote")
def promote_students_cohort(
    payload: StudentPromotionRequest,
    request: Request,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
    current_user: User = Depends(get_current_user),
):
    """
    Promotes a batch of students to the next Class / Division / Academic Year.
    Stores previous class pointers to allow single-click rollback/downgrade if mistakes occur.
    Preserves all registered facial encodings and historical attendance records intact.
    """
    check_tenant_operational_access(tenant)
    target_class = db.query(ClassModel).filter(ClassModel.tenant_id == tenant.id, ClassModel.id == payload.target_class_id).first()
    if not target_class:
        raise HTTPException(status_code=404, detail="Target class not found.")

    students = db.query(Student).filter(
        Student.tenant_id == tenant.id,
        Student.id.in_(payload.student_ids),
    ).all()

    if not students:
        raise HTTPException(status_code=400, detail="No valid students selected for promotion.")

    promoted_count = 0
    now_ist = get_ist_now()

    for s in students:
        # Save previous state for rollback
        s.previous_class_id = s.class_id
        s.previous_division_id = s.division_id
        s.previous_academic_year_id = s.academic_year_id
        s.last_promoted_at = now_ist

        # Apply new academic progression
        s.class_id = payload.target_class_id
        s.division_id = payload.target_division_id
        if payload.target_academic_year_id:
            s.academic_year_id = payload.target_academic_year_id
        s.class_semester = target_class.name
        promoted_count += 1

    # Record Audit Log
    audit = AuditLog(
        tenant_id=tenant.id,
        user_id=current_user.id,
        actor_name=current_user.full_name,
        actor_role=current_user.role,
        action_type="STUDENT_PROMOTION",
        target_type="CLASS",
        target_id=str(target_class.id),
        description=f"Promoted {promoted_count} students to '{target_class.name}'.",
        ip_address=request.client.host if request.client else None,
    )
    db.add(audit)
    db.commit()

    return {
        "status": "success",
        "message": f"Successfully promoted {promoted_count} student(s) to '{target_class.name}'.",
        "promoted_count": promoted_count,
        "target_class": target_class.name,
    }


@router.post("/students/rollback-promotion")
def rollback_student_promotion(
    payload: PromotionRollbackRequest,
    request: Request,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
    current_user: User = Depends(get_current_user),
):
    """
    Rolls back / downgrades a previously promoted batch or entire class back to their previous class state.
    """
    check_tenant_operational_access(tenant)
    query = db.query(Student).filter(Student.tenant_id == tenant.id, Student.previous_class_id.isnot(None))

    if payload.student_ids:
        query = query.filter(Student.id.in_(payload.student_ids))
    elif payload.class_id:
        query = query.filter(Student.class_id == payload.class_id)
    else:
        raise HTTPException(status_code=400, detail="Must provide student_ids or class_id to execute rollback.")

    students = query.all()
    if not students:
        raise HTTPException(status_code=404, detail="No eligible students found with previous promotion history to rollback.")

    rollback_count = 0
    for s in students:
        prev_class_id = s.previous_class_id
        prev_div_id = s.previous_division_id
        prev_acad_id = s.previous_academic_year_id

        # Restore previous state
        prev_class_obj = db.query(ClassModel).filter(ClassModel.id == prev_class_id).first()
        s.class_id = prev_class_id
        s.division_id = prev_div_id
        s.academic_year_id = prev_acad_id
        s.class_semester = prev_class_obj.name if prev_class_obj else "General"

        # Clear previous markers
        s.previous_class_id = None
        s.previous_division_id = None
        s.previous_academic_year_id = None
        rollback_count += 1

    # Record Audit Log
    audit = AuditLog(
        tenant_id=tenant.id,
        user_id=current_user.id,
        actor_name=current_user.full_name,
        actor_role=current_user.role,
        action_type="PROMOTION_ROLLBACK",
        target_type="STUDENT_BATCH",
        target_id=str(rollback_count),
        description=f"Rolled back/downgraded promotion for {rollback_count} student(s) to previous academic cohort.",
        ip_address=request.client.host if request.client else None,
    )
    db.add(audit)
    db.commit()

    return {
        "status": "success",
        "message": f"Successfully rolled back/downgraded {rollback_count} student(s) to their prior classes.",
        "rollback_count": rollback_count,
    }


# ==========================================================
# 8. INDIVIDUAL STUDENT TRANSFER & AUDIT ROLLBACK
# ==========================================================
@router.post("/students/transfer")
def transfer_individual_student(
    payload: StudentTransferRequest,
    request: Request,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
    current_user: User = Depends(get_current_user),
):
    """
    Transfers an individual student to a new Department, Class, and Division.
    Tracks previous assignments with timestamps for audit and rollback.
    Preserves all facial encodings and attendance records intact.
    """
    check_tenant_operational_access(tenant)
    student = db.query(Student).filter(
        Student.tenant_id == tenant.id,
        Student.id == payload.student_id,
        Student.is_active == True,
    ).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found.")

    target_dept = db.query(Department).filter(
        Department.tenant_id == tenant.id,
        Department.id == payload.target_department_id,
    ).first()
    if not target_dept:
        raise HTTPException(status_code=404, detail="Target department not found.")

    target_class = db.query(ClassModel).filter(
        ClassModel.tenant_id == tenant.id,
        ClassModel.id == payload.target_class_id,
    ).first()
    if not target_class:
        raise HTTPException(status_code=404, detail="Target class not found.")

    target_div_name = None
    if payload.target_division_id:
        target_div = db.query(Division).filter(
            Division.tenant_id == tenant.id,
            Division.id == payload.target_division_id,
            Division.class_id == payload.target_class_id,
        ).first()
        if not target_div:
            raise HTTPException(status_code=400, detail="Target division does not belong to the selected class.")
        target_div_name = target_div.name

    # Save previous state for audit & rollback
    now_ist = get_ist_now()
    student.previous_department_id = student.department_id
    student.previous_class_id = student.class_id
    student.previous_division_id = student.division_id
    student.previous_academic_year_id = student.academic_year_id
    student.last_transferred_at = now_ist

    # Apply new transfer destination
    student.department_id = target_dept.id
    student.department = target_dept.name
    student.class_id = target_class.id
    student.class_semester = target_class.name
    student.division_id = payload.target_division_id
    if payload.target_academic_year_id:
        student.academic_year_id = payload.target_academic_year_id

    # Record Audit Log
    audit = AuditLog(
        tenant_id=tenant.id,
        user_id=current_user.id,
        actor_name=current_user.full_name,
        actor_role=current_user.role,
        action_type="STUDENT_TRANSFER",
        target_type="STUDENT",
        target_id=str(student.id),
        description=(
            f"Transferred {student.name} ({student.roll_number}) to Dept: '{target_dept.name}', "
            f"Class: '{target_class.name}'" + (f", Div: '{target_div_name}'" if target_div_name else "") + "."
        ),
        ip_address=request.client.host if request.client else None,
    )
    db.add(audit)
    db.commit()
    db.refresh(student)

    return {
        "status": "success",
        "message": f"Student '{student.name}' transferred to {target_dept.name} - {target_class.name} successfully.",
        "student": student.to_dict(),
    }


@router.post("/students/rollback-transfer")
def rollback_student_transfer(
    payload: StudentTransferRollbackRequest,
    request: Request,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
    current_user: User = Depends(get_current_user),
):
    """
    Reverts an individual student's last transfer back to their previous department, class, and division.
    """
    check_tenant_operational_access(tenant)
    student = db.query(Student).filter(
        Student.tenant_id == tenant.id,
        Student.id == payload.student_id,
        Student.is_active == True,
    ).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found.")

    if not student.previous_department_id and not student.previous_class_id and not student.last_transferred_at:
        raise HTTPException(status_code=400, detail="No prior transfer history found for this student.")

    prev_dept_id = student.previous_department_id
    prev_class_id = student.previous_class_id
    prev_div_id = student.previous_division_id
    prev_year_id = student.previous_academic_year_id

    prev_dept = db.query(Department).filter(Department.id == prev_dept_id).first() if prev_dept_id else None
    prev_class = db.query(ClassModel).filter(ClassModel.id == prev_class_id).first() if prev_class_id else None

    # Revert to previous values
    if prev_dept:
        student.department_id = prev_dept.id
        student.department = prev_dept.name
    elif prev_dept_id:
        student.department_id = prev_dept_id

    if prev_class:
        student.class_id = prev_class.id
        student.class_semester = prev_class.name
    elif prev_class_id:
        student.class_id = prev_class_id

    student.division_id = prev_div_id
    student.academic_year_id = prev_year_id

    # Clear previous markers
    student.previous_department_id = None
    student.previous_class_id = None
    student.previous_division_id = None
    student.previous_academic_year_id = None
    student.last_transferred_at = None

    # Record Audit Log
    audit = AuditLog(
        tenant_id=tenant.id,
        user_id=current_user.id,
        actor_name=current_user.full_name,
        actor_role=current_user.role,
        action_type="TRANSFER_ROLLBACK",
        target_type="STUDENT",
        target_id=str(student.id),
        description=f"Reverted transfer for {student.name} ({student.roll_number}) to previous department & class.",
        ip_address=request.client.host if request.client else None,
    )
    db.add(audit)
    db.commit()
    db.refresh(student)

    return {
        "status": "success",
        "message": f"Successfully reverted transfer for '{student.name}'.",
        "student": student.to_dict(),
    }


@router.get("/transfers/recent")
def list_recent_transfers(
    limit: int = 50,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """
    Returns list of students who have been transferred with previous department/class pointers.
    """
    students = (
        db.query(Student)
        .filter(
            Student.tenant_id == tenant.id,
            Student.is_active == True,
            Student.last_transferred_at.isnot(None),
        )
        .order_by(Student.last_transferred_at.desc())
        .limit(limit)
        .all()
    )

    results = []
    for s in students:
        s_dict = s.to_dict()
        prev_dept = db.query(Department).filter(Department.id == s.previous_department_id).first() if s.previous_department_id else None
        prev_class = db.query(ClassModel).filter(ClassModel.id == s.previous_class_id).first() if s.previous_class_id else None
        prev_div = db.query(Division).filter(Division.id == s.previous_division_id).first() if s.previous_division_id else None
        s_dict["previous_department_name"] = prev_dept.name if prev_dept else "N/A"
        s_dict["previous_class_name"] = prev_class.name if prev_class else "N/A"
        s_dict["previous_division_name"] = prev_div.name if prev_div else "N/A"
        results.append(s_dict)

    return {"status": "success", "transfers": results, "count": len(results)}


# ==========================================================
# 9. EXCEL BULK UPLOAD ENGINE & TEMPLATE GENERATOR
# ==========================================================
@router.get("/students/template")
def download_student_import_template(
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """
    Generates a professionally formatted Excel (.xlsx) workbook with:
    Sheet 1: Student Upload Template (headers, sample rows, styling)
    Sheet 2: Valid Master Reference (active departments, classes, divisions for this tenant)
    """
    wb = openpyxl.Workbook()

    # --- Sheet 1: Template ---
    ws1 = wb.active
    ws1.title = "Student Upload Template"
    ws1.views.sheetView[0].showGridLines = True

    headers = [
        "Full Name *",
        "Roll Number *",
        "Department *",
        "Class / Grade *",
        "Division",
        "Gender (Male/Female/Other)",
        "Role (student/teacher/admin_staff/other)",
        "Email Address",
        "Phone Number",
    ]

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    ws1.row_dimensions[1].height = 28
    for col_num, header_title in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Fetch tenant master records for sample rows
    dept_sample = db.query(Department).filter(Department.tenant_id == tenant.id).order_by(Department.name.asc()).first()
    sample_dept_name = dept_sample.name if dept_sample else "Computer Science"
    
    class_sample = None
    if dept_sample:
        class_sample = db.query(ClassModel).filter(ClassModel.tenant_id == tenant.id, ClassModel.department_id == dept_sample.id).first()
    if not class_sample:
        class_sample = db.query(ClassModel).filter(ClassModel.tenant_id == tenant.id).first()
    sample_class_name = class_sample.name if class_sample else "FY Computer Science"

    div_sample = None
    if class_sample:
        div_sample = db.query(Division).filter(Division.tenant_id == tenant.id, Division.class_id == class_sample.id).first()
    sample_div_name = div_sample.name if div_sample else "Division A"

    sample_rows = [
        [
            "Alexander Hayes",
            "CS-2026-001",
            sample_dept_name,
            sample_class_name,
            sample_div_name,
            "Male",
            "student",
            "alex.hayes@campus.edu",
            "+1234567890",
        ],
        [
            "Sophia Loren",
            "CS-2026-002",
            sample_dept_name,
            sample_class_name,
            sample_div_name,
            "Female",
            "student",
            "sophia.loren@campus.edu",
            "+1987654321",
        ],
        [
            "Jordan Reed",
            "CS-2026-003",
            sample_dept_name,
            sample_class_name,
            "",
            "Other",
            "student",
            "jordan.reed@campus.edu",
            "",
        ],
    ]

    data_font = Font(name="Segoe UI", size=10)
    data_align = Alignment(vertical="center")

    for row_idx, row_data in enumerate(sample_rows, start=2):
        ws1.row_dimensions[row_idx].height = 22
        for col_idx, val in enumerate(row_data, start=1):
            c = ws1.cell(row=row_idx, column=col_idx, value=val)
            c.font = data_font
            c.alignment = data_align
            c.border = thin_border

    # Set column widths
    col_widths = [24, 18, 28, 24, 16, 26, 26, 28, 18]
    for i, w in enumerate(col_widths, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # --- Sheet 2: Master Reference Data ---
    ws2 = wb.create_sheet(title="Valid Master Reference")
    ws2.views.sheetView[0].showGridLines = True

    ref_headers = [
        "Valid Department Name",
        "Department Code",
        "Valid Class Name",
        "Class Parent Department",
        "Valid Division Name",
        "Division Parent Class",
    ]
    ref_header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    ref_header_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")

    ws2.row_dimensions[1].height = 28
    for col_idx, h in enumerate(ref_headers, 1):
        c = ws2.cell(row=1, column=col_idx, value=h)
        c.font = ref_header_font
        c.fill = ref_header_fill
        c.alignment = header_align
        c.border = thin_border

    depts = db.query(Department).filter(Department.tenant_id == tenant.id).order_by(Department.name.asc()).all()
    classes = db.query(ClassModel).filter(ClassModel.tenant_id == tenant.id).order_by(ClassModel.name.asc()).all()
    divs = db.query(Division).filter(Division.tenant_id == tenant.id).order_by(Division.name.asc()).all()

    max_len = max(len(depts), len(classes), len(divs), 1)
    for row_idx in range(max_len):
        r = row_idx + 2
        ws2.row_dimensions[r].height = 20
        d = depts[row_idx] if row_idx < len(depts) else None
        cl = classes[row_idx] if row_idx < len(classes) else None
        dv = divs[row_idx] if row_idx < len(divs) else None

        c1 = ws2.cell(row=r, column=1, value=d.name if d else "")
        c2 = ws2.cell(row=r, column=2, value=d.code if d else "")
        c3 = ws2.cell(row=r, column=3, value=cl.name if cl else "")
        c4 = ws2.cell(row=r, column=4, value=cl.department if cl else "")
        c5 = ws2.cell(row=r, column=5, value=dv.name if dv else "")
        c6 = ws2.cell(row=r, column=6, value=dv.class_obj.name if dv and dv.class_obj else "")

        for col_i in range(1, 7):
            cell = ws2.cell(row=r, column=col_i)
            cell.font = data_font
            cell.alignment = data_align
            cell.border = thin_border

    ref_widths = [28, 16, 26, 26, 20, 26]
    for i, w in enumerate(ref_widths, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Student_Bulk_Upload_Template_{tenant.slug or tenant.id}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/students/bulk-preview")
async def preview_student_bulk_upload(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """
    Parses and validates an uploaded Excel file against tenant master tables.
    Flags invalid rows with exact reasons (Option A: partial staging allowed).
    Rejects duplicate roll numbers in file and existing roll numbers in tenant DB.
    """
    check_tenant_operational_access(tenant)

    if not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Invalid file format. Please upload an Excel workbook (.xlsx or .xls).")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel workbook: {str(e)}")

    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise HTTPException(status_code=400, detail="The uploaded sheet is empty or contains no data rows.")

    # Parse headers
    header_row = [str(c).strip().lower() if c is not None else "" for c in rows[0]]

    def find_col_idx(candidates: List[str]) -> Optional[int]:
        for cand in candidates:
            for idx, h in enumerate(header_row):
                if cand in h:
                    return idx
        return None

    name_idx = find_col_idx(["full name", "student name", "name", "full_name"])
    roll_idx = find_col_idx(["roll number", "roll no", "roll_no", "roll", "student id", "id"])
    dept_idx = find_col_idx(["department", "dept"])
    class_idx = find_col_idx(["class / grade", "class name", "grade", "class", "standard"])
    div_idx = find_col_idx(["division", "section", "div"])
    gender_idx = find_col_idx(["gender", "sex"])
    role_idx = find_col_idx(["role", "user role", "user_role", "type"])
    email_idx = find_col_idx(["email", "mail"])
    phone_idx = find_col_idx(["phone", "mobile", "contact"])

    if name_idx is None or roll_idx is None or dept_idx is None or class_idx is None:
        raise HTTPException(
            status_code=400,
            detail="Required column headers missing. Ensure 'Full Name', 'Roll Number', 'Department', and 'Class / Grade' columns exist.",
        )

    # Fetch Master Lookups
    departments = db.query(Department).filter(Department.tenant_id == tenant.id).all()
    dept_by_name = {d.name.strip().lower(): d for d in departments}
    dept_by_code = {d.code.strip().lower(): d for d in departments if d.code}

    classes = db.query(ClassModel).filter(ClassModel.tenant_id == tenant.id).all()
    classes_by_name = {}
    for c in classes:
        key = c.name.strip().lower()
        classes_by_name.setdefault(key, []).append(c)

    divisions = db.query(Division).filter(Division.tenant_id == tenant.id).all()
    div_lookup = {}
    for dv in divisions:
        div_lookup[(dv.class_id, dv.name.strip().lower())] = dv

    # Existing tenant students roll numbers
    existing_rolls = {
        r[0].strip().lower()
        for r in db.query(Student.roll_number).filter(Student.tenant_id == tenant.id, Student.is_active == True).all()
        if r[0]
    }

    seen_file_rolls = set()
    preview_items = []
    valid_count = 0
    invalid_count = 0

    for row_num, row in enumerate(rows[1:], start=2):
        # Skip completely empty rows
        if not any(row):
            continue

        raw_name = str(row[name_idx]).strip() if name_idx < len(row) and row[name_idx] is not None else ""
        raw_roll = str(row[roll_idx]).strip() if roll_idx < len(row) and row[roll_idx] is not None else ""
        raw_dept = str(row[dept_idx]).strip() if dept_idx < len(row) and row[dept_idx] is not None else ""
        raw_class = str(row[class_idx]).strip() if class_idx < len(row) and row[class_idx] is not None else ""
        raw_div = str(row[div_idx]).strip() if div_idx is not None and div_idx < len(row) and row[div_idx] is not None else ""
        raw_gender = str(row[gender_idx]).strip() if gender_idx is not None and gender_idx < len(row) and row[gender_idx] is not None else "Other"
        raw_role = str(row[role_idx]).strip().lower() if role_idx is not None and role_idx < len(row) and row[role_idx] is not None else "student"
        raw_email = str(row[email_idx]).strip() if email_idx is not None and email_idx < len(row) and row[email_idx] is not None else ""
        raw_phone = str(row[phone_idx]).strip() if phone_idx is not None and phone_idx < len(row) and row[phone_idx] is not None else ""

        errors = []
        resolved_dept = None
        resolved_class = None
        resolved_div = None

        # 1. Validate Name
        if not raw_name:
            errors.append("Full Name cannot be empty.")

        # 2. Validate Roll Number
        if not raw_roll:
            errors.append("Roll Number cannot be empty.")
        else:
            roll_lower = raw_roll.lower()
            if roll_lower in seen_file_rolls:
                errors.append(f"Duplicate Roll Number '{raw_roll}' found in upload file.")
            elif roll_lower in existing_rolls:
                errors.append(f"Roll Number '{raw_roll}' is already registered in this institution.")
            else:
                seen_file_rolls.add(roll_lower)

        # 3. Validate Department
        if not raw_dept:
            errors.append("Department cannot be empty.")
        else:
            dept_key = raw_dept.lower()
            resolved_dept = dept_by_name.get(dept_key) or dept_by_code.get(dept_key)
            if not resolved_dept:
                errors.append(f"Department '{raw_dept}' does not exist in institution master.")

        # 4. Validate Class
        if not raw_class:
            errors.append("Class / Grade cannot be empty.")
        else:
            class_key = raw_class.lower()
            matching_classes = classes_by_name.get(class_key, [])
            if not matching_classes:
                errors.append(f"Class '{raw_class}' does not exist in institution master.")
            else:
                if resolved_dept:
                    dept_matched = [c for c in matching_classes if c.department_id == resolved_dept.id or c.department.lower() == resolved_dept.name.lower()]
                    if dept_matched:
                        resolved_class = dept_matched[0]
                    else:
                        resolved_class = matching_classes[0]
                        if resolved_class.department_id and resolved_class.department_id != resolved_dept.id:
                            errors.append(f"Class '{raw_class}' belongs to '{resolved_class.department}', not '{resolved_dept.name}'.")
                else:
                    resolved_class = matching_classes[0]

        # 5. Validate Division (optional)
        if raw_div and resolved_class:
            div_key = (resolved_class.id, raw_div.lower())
            resolved_div = div_lookup.get(div_key)
            if not resolved_div:
                div_matched = [dv for dv in divisions if dv.class_id == resolved_class.id and dv.name.lower() == raw_div.lower()]
                if div_matched:
                    resolved_div = div_matched[0]
                else:
                    errors.append(f"Division '{raw_div}' is not registered under class '{resolved_class.name}'.")

        # 6. Validate Gender
        gender_clean = raw_gender.capitalize()
        if gender_clean not in ["Male", "Female", "Other"]:
            gender_clean = "Other"

        # 7. Validate Role
        if raw_role not in ["student", "teacher", "admin_staff", "other"]:
            raw_role = "student"

        is_valid = len(errors) == 0
        if is_valid:
            valid_count += 1
        else:
            invalid_count += 1

        preview_items.append({
            "row_index": row_num,
            "name": raw_name,
            "roll_number": raw_roll,
            "department": resolved_dept.name if resolved_dept else raw_dept,
            "department_id": resolved_dept.id if resolved_dept else None,
            "class_name": resolved_class.name if resolved_class else raw_class,
            "class_id": resolved_class.id if resolved_class else None,
            "division_name": resolved_div.name if resolved_div else (raw_div or "N/A"),
            "division_id": resolved_div.id if resolved_div else None,
            "gender": gender_clean,
            "user_role": raw_role,
            "email": raw_email,
            "phone_number": raw_phone,
            "is_valid": is_valid,
            "errors": errors,
        })

    return {
        "status": "success",
        "filename": file.filename,
        "total_rows": len(preview_items),
        "valid_count": valid_count,
        "invalid_count": invalid_count,
        "preview_data": preview_items,
    }


@router.post("/students/bulk-import")
def execute_student_bulk_import(
    payload: BulkStudentImportRequest,
    request: Request,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
    current_user: User = Depends(get_current_user),
):
    """
    Executes bulk student creation for validated rows, tags them with a batch_upload_id,
    and records an audit log.
    """
    check_tenant_operational_access(tenant)

    if not payload.rows:
        raise HTTPException(status_code=400, detail="No student rows provided for import.")

    # Create Batch record
    batch = StudentBatchUpload(
        tenant_id=tenant.id,
        filename=payload.filename or "bulk_upload.xlsx",
        uploaded_by_user_id=current_user.id,
        total_rows=len(payload.rows),
        valid_rows=len(payload.rows),
        imported_count=0,
        is_active=True,
    )
    db.add(batch)
    db.flush()

    imported_students = []
    skipped_duplicates = 0

    academic_year_id = payload.academic_year_id
    if not academic_year_id:
        curr_year = db.query(AcademicYear).filter(AcademicYear.tenant_id == tenant.id, AcademicYear.is_current == True).first()
        if curr_year:
            academic_year_id = curr_year.id

    for item in payload.rows:
        exists = db.query(Student).filter(
            Student.tenant_id == tenant.id,
            Student.roll_number == item.roll_number.strip(),
            Student.is_active == True,
        ).first()
        if exists:
            skipped_duplicates += 1
            continue

        dept = db.query(Department).filter(Department.tenant_id == tenant.id, Department.id == item.department_id).first()
        class_obj = db.query(ClassModel).filter(ClassModel.tenant_id == tenant.id, ClassModel.id == item.class_id).first()

        dept_name = dept.name if dept else "Computer Science"
        class_name = class_obj.name if class_obj else "General"

        student = Student(
            tenant_id=tenant.id,
            name=item.name.strip(),
            roll_number=item.roll_number.strip(),
            department_id=item.department_id,
            department=dept_name,
            class_id=item.class_id,
            class_semester=class_name,
            division_id=item.division_id,
            academic_year_id=academic_year_id,
            gender=item.gender or "Other",
            user_role=item.user_role or "student",
            email=item.email.strip() if item.email else None,
            phone_number=item.phone_number.strip() if item.phone_number else None,
            batch_upload_id=batch.id,
            is_active=True,
        )
        db.add(student)
        imported_students.append(student)

    batch.imported_count = len(imported_students)

    # Record Audit Log
    audit = AuditLog(
        tenant_id=tenant.id,
        user_id=current_user.id,
        actor_name=current_user.full_name,
        actor_role=current_user.role,
        action_type="BULK_STUDENT_IMPORT",
        target_type="STUDENT_BATCH",
        target_id=str(batch.id),
        description=f"Imported {batch.imported_count} student(s) from batch '{batch.filename}' (Batch #{batch.id}).",
        ip_address=request.client.host if request.client else None,
    )
    db.add(audit)
    db.commit()
    db.refresh(batch)

    return {
        "status": "success",
        "message": f"Successfully imported {batch.imported_count} student(s) under Batch #{batch.id}.",
        "batch": batch.to_dict(),
        "imported_count": batch.imported_count,
        "skipped_duplicates": skipped_duplicates,
    }


# ==========================================================
# 10. BATCH UPLOAD AUDIT & BATCH SOFT-DELETE
# ==========================================================
@router.get("/batches")
def list_student_batches(
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
):
    """
    Returns history of student batch uploads with active student count.
    """
    batches = (
        db.query(StudentBatchUpload)
        .filter(StudentBatchUpload.tenant_id == tenant.id)
        .order_by(StudentBatchUpload.id.desc())
        .all()
    )

    results = []
    for b in batches:
        b_dict = b.to_dict()
        active_count = (
            db.query(Student)
            .filter(
                Student.tenant_id == tenant.id,
                Student.batch_upload_id == b.id,
                Student.is_active == True,
            )
            .count()
        )
        b_dict["active_students_count"] = active_count
        results.append(b_dict)

    return {"status": "success", "batches": results, "count": len(results)}


@router.delete("/batches/{batch_id}")
def delete_student_batch(
    batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_current_tenant),
    current_user: User = Depends(get_current_user),
):
    """
    Soft-deletes an entire bulk upload batch (is_active = False) and all associated students.
    Immediately hot-reloads FaceEngine live vector RAM cache so soft-deleted students
    are excluded from live facial recognition.
    """
    check_tenant_operational_access(tenant)

    batch = db.query(StudentBatchUpload).filter(
        StudentBatchUpload.tenant_id == tenant.id,
        StudentBatchUpload.id == batch_id,
    ).first()
    if not batch:
        raise HTTPException(status_code=404, detail="Batch upload record not found.")

    # Soft delete the batch record
    batch.is_active = False

    # Soft delete all active students associated with this batch
    soft_deleted_count = (
        db.query(Student)
        .filter(
            Student.tenant_id == tenant.id,
            Student.batch_upload_id == batch_id,
            Student.is_active == True,
        )
        .update({"is_active": False})
    )

    # Hot-reload FaceEngine in-memory RAM cache for this tenant
    try:
        face_engine.reload_cache(db, tenant_id=tenant.id)
        logger.info(f"FaceEngine cache reloaded after soft-deleting batch {batch_id} for tenant {tenant.id}.")
    except Exception as e:
        logger.error(f"Failed to hot-reload FaceEngine cache for tenant {tenant.id}: {e}")

    # Record Audit Log
    audit = AuditLog(
        tenant_id=tenant.id,
        user_id=current_user.id,
        actor_name=current_user.full_name,
        actor_role=current_user.role,
        action_type="BATCH_SOFT_DELETE",
        target_type="STUDENT_BATCH",
        target_id=str(batch.id),
        description=(
            f"Soft-deleted batch #{batch.id} ('{batch.filename}'), deactivating {soft_deleted_count} student(s) "
            f"and hot-reloading live biometric RAM cache."
        ),
        ip_address=request.client.host if request.client else None,
    )
    db.add(audit)
    db.commit()

    return {
        "status": "success",
        "message": f"Batch #{batch_id} and {soft_deleted_count} student(s) soft-deleted successfully. Biometric RAM cache hot-reloaded.",
        "batch_id": batch_id,
        "soft_deleted_students": soft_deleted_count,
    }
