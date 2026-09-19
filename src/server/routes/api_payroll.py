import io
import json
import logging
from datetime import datetime, date, timedelta, time as dt_time
from typing import Optional, List, Dict, Any
from fastapi import APIRouter, Depends, Query, HTTPException, Request, status
from fastapi.responses import StreamingResponse, JSONResponse
from pydantic import BaseModel, Field
import pandas as pd
from sqlalchemy.orm import Session
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from src.database.models import (
    Student,
    AttendanceRecord,
    SystemBranding,
    Tenant,
    Department,
    User,
    LeaveRequest,
    LeaveType,
    CompanyLocation,
    DesignationMaster,
    SalaryComponent,
    SalaryTemplate,
    EmployeeSalaryStructure,
    SalaryRevisionHistory,
    PayrollBatch,
    PayrollPayslip,
)
from src.database.session import get_db
from src.server.tenant_middleware import get_current_tenant, resolve_tenant
from src.server.rbac_middleware import require_roles, check_tenant_operational_access, check_tenant_payroll_access, get_current_user
from src.utils.timezone import get_ist_now, get_ist_date
from src.core.payroll_engine import (
    calculate_employee_payroll,
    number_to_words_inr,
    get_effective_salary_structures,
)
from src.services.payroll_service import PayrollService
from src.services.organization_service import OrganizationService

logger = logging.getLogger("api_payroll")
router = APIRouter(prefix="/api/v1/payroll", tags=["Payroll & Compensation"])


def get_current_payroll_tenant(tenant: Tenant = Depends(get_current_tenant)) -> Tenant:
    check_tenant_payroll_access(tenant)
    return tenant



# ==============================================================================
# PYDANTIC REQUEST SCHEMAS
# ==============================================================================

class LocationCreateRequest(BaseModel):
    name: str
    code: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = "Maharashtra"
    address: Optional[str] = None
    contact_number: Optional[str] = None
    is_active: Optional[bool] = True


class LocationUpdateRequest(BaseModel):
    name: Optional[str] = None
    code: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    address: Optional[str] = None
    contact_number: Optional[str] = None
    is_active: Optional[bool] = None


class DesignationCreateRequest(BaseModel):
    title: str
    code: Optional[str] = None
    department_id: Optional[int] = None
    salary_template_id: Optional[int] = None
    description: Optional[str] = None
    is_active: Optional[bool] = True


class DesignationUpdateRequest(BaseModel):
    title: Optional[str] = None
    code: Optional[str] = None
    department_id: Optional[int] = None
    salary_template_id: Optional[int] = None
    description: Optional[str] = None
    is_active: Optional[bool] = None


class SalaryComponentCreateRequest(BaseModel):
    name: str
    code: str
    component_type: str = "EARNING"
    calculation_type: str = "FIXED"
    default_value: float = 0.0
    is_taxable: bool = True
    is_statutory: bool = False
    is_active: bool = True


class SalaryComponentUpdateRequest(BaseModel):
    name: Optional[str] = None
    component_type: Optional[str] = None
    calculation_type: Optional[str] = None
    default_value: Optional[float] = None
    is_taxable: Optional[bool] = None
    is_statutory: Optional[bool] = None
    is_active: Optional[bool] = None


class SalaryTemplateCreateRequest(BaseModel):
    name: str
    code: str
    compensation_model: str = "STRUCTURED_SALARY"
    description: Optional[str] = None
    basic_percentage: float = 50.0
    hra_percentage: float = 20.0
    da_percentage: float = 0.0
    conveyance_fixed: float = 1600.0
    medical_fixed: float = 1250.0
    enable_pf: bool = True
    pf_capped_at_ceiling: bool = True
    enable_esi: bool = True
    enable_pt: bool = True
    is_active: bool = True


class SalaryTemplateUpdateRequest(BaseModel):
    name: Optional[str] = None
    compensation_model: Optional[str] = None
    description: Optional[str] = None
    basic_percentage: Optional[float] = None
    hra_percentage: Optional[float] = None
    da_percentage: Optional[float] = None
    conveyance_fixed: Optional[float] = None
    medical_fixed: Optional[float] = None
    enable_pf: Optional[bool] = None
    pf_capped_at_ceiling: Optional[bool] = None
    enable_esi: Optional[bool] = None
    enable_pt: Optional[bool] = None
    is_active: Optional[bool] = None


class EmployeeStructureAssignRequest(BaseModel):
    student_id: int
    template_id: Optional[int] = None
    compensation_model: str = "STRUCTURED_SALARY"
    annual_ctc: Optional[float] = 0.0
    monthly_gross: Optional[float] = 0.0
    monthly_basic: Optional[float] = 0.0
    monthly_da: Optional[float] = 0.0
    monthly_hra: Optional[float] = 0.0
    conveyance_allowance: Optional[float] = 0.0
    medical_allowance: Optional[float] = 0.0
    special_allowance: Optional[float] = 0.0
    other_allowances: Optional[float] = 0.0
    hourly_rate: Optional[float] = 0.0
    daily_rate: Optional[float] = 0.0
    fixed_stipend: Optional[float] = 0.0
    commission_percentage: Optional[float] = 0.0
    enable_pf: Optional[bool] = True
    pf_capped_at_ceiling: Optional[bool] = True
    enable_esi: Optional[bool] = True
    enable_pt: Optional[bool] = True
    pt_monthly_amount: Optional[float] = 200.0
    tds_monthly_amount: Optional[float] = 0.0
    effective_from_date: str = Field(..., description="YYYY-MM-DD")
    revision_reason: Optional[str] = "Salary Revision / Placement"


class EmployeeStatutoryBankingRequest(BaseModel):
    student_id: int
    pan_number: Optional[str] = None
    uan_number: Optional[str] = None
    esic_number: Optional[str] = None
    bank_name: Optional[str] = None
    bank_account_number: Optional[str] = None
    bank_ifsc_code: Optional[str] = None
    location_id: Optional[int] = None
    designation_id: Optional[int] = None


class BatchGenerateRequest(BaseModel):
    period_year: int
    period_month: int
    working_days: Optional[float] = 26.0
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    department_id: Optional[int] = None


class BatchStatusUpdateRequest(BaseModel):
    remarks: Optional[str] = None


class EmployeeRateUpdateRequest(BaseModel):
    student_id: int
    hourly_rate: Optional[float] = None
    monthly_base_salary: Optional[float] = None


class PayrollSettingsUpdateRequest(BaseModel):
    payroll_structure: Optional[str] = "STRUCTURED_SALARY"
    default_hourly_rate: Optional[float] = 15.0
    standard_working_hours_per_day: Optional[float] = 8.0
    enable_overtime: Optional[bool] = True
    overtime_rate_multiplier: Optional[float] = 1.5
    holiday_ot_multiplier: Optional[float] = 2.0
    missed_checkout_policy: Optional[str] = "HALF_DAY"
    currency_symbol: Optional[str] = "₹"
    epf_employee_pct: Optional[float] = 12.0
    epf_employer_pct: Optional[float] = 12.0
    epf_wage_ceiling: Optional[float] = 15000.0
    epf_admin_charges_pct: Optional[float] = 0.50
    esi_employee_pct: Optional[float] = 0.75
    esi_employer_pct: Optional[float] = 3.25
    esi_gross_threshold: Optional[float] = 21000.0
    pt_monthly_default: Optional[float] = 200.0
    enable_pf_ceiling: Optional[bool] = True


# ==============================================================================
# 1. ORGANIZATION MASTERS ENDPOINTS
# ==============================================================================

# --- A) Locations ---
@router.get("/masters/locations")
def list_locations(
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return {"status": "success", "data": org_service.list_locations(current_tenant.id)}


@router.get("/masters/locations/{loc_id}")
def get_location(
    loc_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return {"status": "success", "data": org_service.get_location(current_tenant.id, loc_id)}


@router.post("/masters/locations")
def create_location(
    payload: LocationCreateRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    data = org_service.create_location(current_tenant.id, payload)
    return {"status": "success", "message": "Location created successfully.", "data": data}


@router.put("/masters/locations/{loc_id}")
def update_location(
    loc_id: int,
    payload: LocationUpdateRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    data = org_service.update_location(current_tenant.id, loc_id, payload)
    return {"status": "success", "message": "Location updated successfully.", "data": data}


@router.delete("/masters/locations/{loc_id}")
def delete_location(
    loc_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return org_service.delete_location(current_tenant.id, loc_id)


# --- B) Designations ---
@router.get("/masters/designations")
def list_designations(
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return {"status": "success", "data": org_service.list_designations(current_tenant.id)}


@router.get("/masters/designations/{desig_id}")
def get_designation(
    desig_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return {"status": "success", "data": org_service.get_designation(current_tenant.id, desig_id)}


@router.post("/masters/designations")
def create_designation(
    payload: DesignationCreateRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    data = org_service.create_designation(current_tenant.id, payload)
    return {"status": "success", "message": "Designation created.", "data": data}


@router.put("/masters/designations/{desig_id}")
def update_designation(
    desig_id: int,
    payload: DesignationUpdateRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    data = org_service.update_designation(current_tenant.id, desig_id, payload)
    return {"status": "success", "message": "Designation updated.", "data": data}


@router.delete("/masters/designations/{desig_id}")
def delete_designation(
    desig_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return org_service.delete_designation(current_tenant.id, desig_id)


# --- C) Salary Components ---
@router.get("/masters/components")
def list_salary_components(
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return {"status": "success", "data": org_service.list_salary_components(current_tenant.id)}


@router.get("/masters/components/{comp_id}")
def get_salary_component(
    comp_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return {"status": "success", "data": org_service.get_salary_component(current_tenant.id, comp_id)}


@router.post("/masters/components")
def create_salary_component(
    payload: SalaryComponentCreateRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    data = org_service.create_salary_component(current_tenant.id, payload)
    return {"status": "success", "message": "Salary component created.", "data": data}


@router.put("/masters/components/{comp_id}")
def update_salary_component(
    comp_id: int,
    payload: SalaryComponentUpdateRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    data = org_service.update_salary_component(current_tenant.id, comp_id, payload)
    return {"status": "success", "message": "Salary component updated.", "data": data}


@router.delete("/masters/components/{comp_id}")
def delete_salary_component(
    comp_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return org_service.delete_salary_component(current_tenant.id, comp_id)


# --- D) Salary Templates ---
@router.get("/masters/templates")
def list_salary_templates(
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return {"status": "success", "data": org_service.list_salary_templates(current_tenant.id)}


@router.get("/masters/templates/{tpl_id}")
def get_salary_template(
    tpl_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return {"status": "success", "data": org_service.get_salary_template(current_tenant.id, tpl_id)}


@router.post("/masters/templates")
def create_salary_template(
    payload: SalaryTemplateCreateRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    data = org_service.create_salary_template(current_tenant.id, payload)
    return {"status": "success", "message": "Salary template created.", "data": data}


@router.put("/masters/templates/{tpl_id}")
def update_salary_template(
    tpl_id: int,
    payload: SalaryTemplateUpdateRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    data = org_service.update_salary_template(current_tenant.id, tpl_id, payload)
    return {"status": "success", "message": "Salary template updated.", "data": data}


@router.delete("/masters/templates/{tpl_id}")
def delete_salary_template(
    tpl_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    org_service = OrganizationService(db)
    return org_service.delete_salary_template(current_tenant.id, tpl_id)


# ==============================================================================
# 2. EMPLOYEE SALARY STRUCTURES & STATUTORY REVISION
# ==============================================================================

@router.get("/structures")
@router.get("/employees")
def list_employee_structures(
    department_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return {"status": "success", "data": payroll_service.get_structure_overview(current_tenant, department_id)}


@router.get("/employee/{student_id}/structure")
@router.get("/employees/{student_id}/structure")
@router.get("/structures/{student_id}")
def get_employee_structure(
    student_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return payroll_service.get_employee_structure(current_tenant.id, student_id)


@router.post("/structure/assign")
def assign_employee_salary_structure(
    payload: EmployeeStructureAssignRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
    current_user: User = Depends(get_current_user),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return payroll_service.assign_salary_structure(current_tenant.id, payload, current_user)


@router.post("/employee/{student_id}/statutory-banking")
def update_employee_statutory_banking(
    student_id: int,
    payload: EmployeeStatutoryBankingRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return payroll_service.update_statutory_and_banking(current_tenant.id, student_id, payload)


# ==============================================================================
# 3. PAYROLL BATCHES & MONTHLY RUN LIFECYCLE
# ==============================================================================

@router.get("/batches")
def list_payroll_batches(
    year: Optional[int] = Query(None),
    status_filter: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return {"status": "success", "data": payroll_service.list_batches(current_tenant.id, year, status_filter)}


@router.post("/batches/generate")
def generate_payroll_batch(
    payload: BatchGenerateRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
    current_user: User = Depends(get_current_user),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return payroll_service.generate_batch(current_tenant, payload, current_user)


@router.get("/batches/{batch_id}")
def get_payroll_batch_details(
    batch_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return payroll_service.get_batch_details(current_tenant.id, batch_id)


@router.post("/batches/{batch_id}/verify")
def verify_payroll_batch(
    batch_id: int,
    payload: BatchStatusUpdateRequest = None,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
    current_user: User = Depends(get_current_user),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return payroll_service.update_batch_status(current_tenant.id, batch_id, "VERIFY", current_user)


@router.post("/batches/{batch_id}/approve")
def approve_payroll_batch(
    batch_id: int,
    payload: BatchStatusUpdateRequest = None,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
    current_user: User = Depends(get_current_user),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return payroll_service.update_batch_status(current_tenant.id, batch_id, "APPROVE", current_user)


@router.post("/batches/{batch_id}/disburse")
def disburse_payroll_batch(
    batch_id: int,
    payload: BatchStatusUpdateRequest = None,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
    current_user: User = Depends(get_current_user),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return payroll_service.update_batch_status(current_tenant.id, batch_id, "DISBURSE", current_user)


@router.delete("/batches/{batch_id}")
def delete_payroll_batch(
    batch_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return payroll_service.delete_batch(current_tenant.id, batch_id)


@router.get("/batches/{batch_id}/export-bank-advice")
def export_bank_disbursement_advice(
    batch_id: int,
    export_format: str = Query("xlsx", pattern="^(csv|xlsx)$"),
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    batch = (
        db.query(PayrollBatch)
        .filter(PayrollBatch.id == batch_id, PayrollBatch.tenant_id == current_tenant.id)
        .first()
    )
    if not batch:
        raise HTTPException(status_code=404, detail="Payroll batch not found.")

    payslips = (
        db.query(PayrollPayslip)
        .filter(PayrollPayslip.batch_id == batch.id, PayrollPayslip.tenant_id == current_tenant.id)
        .all()
    )

    branding = current_tenant.branding
    currency = branding.currency_symbol if branding and branding.currency_symbol else "₹"
    company_name = branding.institution_name if branding else current_tenant.name

    rows = []
    for p in payslips:
        emp = p.student
        rows.append({
            "Beneficiary Code": emp.roll_number if emp else "",
            "Beneficiary Name": emp.name if emp else "",
            "Bank Name": emp.bank_name if emp and emp.bank_name else "N/A",
            "Account Number": emp.bank_account_number if emp and emp.bank_account_number else "N/A",
            "IFSC Code": emp.bank_ifsc_code if emp and emp.bank_ifsc_code else "N/A",
            "Net Payable Amount": p.net_salary,
            "Payment Reference": p.payment_ref_no or f"SAL-{p.period_year}{p.period_month:02d}-{p.id}",
            "Payment Status": p.payment_status,
        })

    df = pd.DataFrame(rows)
    filename_base = f"bank_disbursement_advice_{batch.batch_number}"

    if export_format == "csv":
        csv_buffer = io.StringIO()
        csv_buffer.write(f"# COMPANY: {company_name}\n")
        csv_buffer.write(f"# BATCH: {batch.batch_number} | PERIOD: {batch.period_month}/{batch.period_year}\n")
        csv_buffer.write(f"# TOTAL NET DISBURSEMENT: {currency}{batch.total_net_outlay:.2f}\n")
        df.to_csv(csv_buffer, index=False)
        csv_buffer.seek(0)
        return StreamingResponse(
            io.BytesIO(csv_buffer.getvalue().encode("utf-8")),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename_base}.csv"},
        )
    else:
        excel_buffer = io.BytesIO()
        with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Bank Advice", startrow=4)
            ws = writer.sheets["Bank Advice"]

            ws.merge_cells("A1:H1")
            ws["A1"] = f"{company_name} — Bank Disbursement Advice"
            ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
            ws["A1"].fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells("A2:H2")
            ws["A2"] = (
                f"Batch: {batch.batch_number} | Period: {batch.period_month}/{batch.period_year} | "
                f"Total Net Payout: {currency}{batch.total_net_outlay:,.2f} | Beneficiaries: {len(payslips)}"
            )
            ws["A2"].font = Font(name="Calibri", size=10, italic=True)
            ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        excel_buffer.seek(0)
        return StreamingResponse(
            excel_buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename_base}.xlsx"},
        )


# ==============================================================================
# 4. INDIVIDUAL PAYSLIP ENDPOINT
# ==============================================================================

@router.get("/payslips/{payslip_id}")
def get_payslip_details(
    payslip_id: int,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return payroll_service.get_payslip(current_tenant, payslip_id)


# ==============================================================================
# 5. STATUTORY & ORGANIZATION SETTINGS
# ==============================================================================

@router.get("/settings")
def get_payroll_settings(
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    branding = current_tenant.branding
    if not branding:
        branding = SystemBranding(tenant_id=current_tenant.id)
        db.add(branding)
        db.commit()
        db.refresh(branding)

    return {
        "status": "success",
        "data": {
            "payroll_structure": branding.payroll_structure or "STRUCTURED_SALARY",
            "default_hourly_rate": float(branding.default_hourly_rate or 15.0),
            "standard_working_hours_per_day": float(branding.standard_working_hours_per_day or 8.0),
            "enable_overtime": bool(branding.enable_overtime),
            "overtime_rate_multiplier": float(branding.overtime_rate_multiplier or 1.5),
            "holiday_ot_multiplier": float(branding.holiday_ot_multiplier or 2.0),
            "missed_checkout_policy": branding.missed_checkout_policy or "HALF_DAY",
            "currency_symbol": branding.currency_symbol or "₹",
            "epf_employee_pct": float(branding.epf_employee_pct or 12.0),
            "epf_employer_pct": float(branding.epf_employer_pct or 12.0),
            "epf_wage_ceiling": float(branding.epf_ceiling_limit or 15000.0),
            "epf_admin_charges_pct": float(branding.epf_admin_charges_pct or 0.50),
            "esi_employee_pct": float(branding.esic_employee_pct or 0.75),
            "esi_employer_pct": float(branding.esic_employer_pct or 3.25),
            "esi_gross_threshold": float(branding.esi_gross_threshold or 21000.0),
            "pt_monthly_default": float(branding.pt_monthly_default or 200.0),
            "enable_pf_ceiling": bool(branding.enable_pf_ceiling),
        },
    }


@router.post("/settings")
def update_payroll_settings(
    payload: PayrollSettingsUpdateRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    branding = current_tenant.branding
    if not branding:
        branding = SystemBranding(tenant_id=current_tenant.id)
        db.add(branding)

    if payload.payroll_structure:
        branding.payroll_structure = payload.payroll_structure.strip().upper()
    if payload.default_hourly_rate is not None:
        branding.default_hourly_rate = max(0.0, float(payload.default_hourly_rate))
    if payload.standard_working_hours_per_day is not None:
        branding.standard_working_hours_per_day = max(1.0, float(payload.standard_working_hours_per_day))
    if payload.enable_overtime is not None:
        branding.enable_overtime = bool(payload.enable_overtime)
    if payload.overtime_rate_multiplier is not None:
        branding.overtime_rate_multiplier = max(1.0, float(payload.overtime_rate_multiplier))
    if payload.holiday_ot_multiplier is not None:
        branding.holiday_ot_multiplier = max(1.0, float(payload.holiday_ot_multiplier))
    if payload.missed_checkout_policy:
        branding.missed_checkout_policy = payload.missed_checkout_policy.strip().upper()
    if payload.currency_symbol:
        branding.currency_symbol = payload.currency_symbol.strip()

    if payload.epf_employee_pct is not None:
        branding.epf_employee_pct = max(0.0, float(payload.epf_employee_pct))
    if payload.epf_employer_pct is not None:
        branding.epf_employer_pct = max(0.0, float(payload.epf_employer_pct))
    if payload.epf_wage_ceiling is not None:
        branding.epf_ceiling_limit = max(0.0, float(payload.epf_wage_ceiling))
    if payload.epf_admin_charges_pct is not None:
        branding.epf_admin_charges_pct = max(0.0, float(payload.epf_admin_charges_pct))
    if payload.esi_employee_pct is not None:
        branding.esic_employee_pct = max(0.0, float(payload.esi_employee_pct))
    if payload.esi_employer_pct is not None:
        branding.esic_employer_pct = max(0.0, float(payload.esi_employer_pct))
    if payload.esi_gross_threshold is not None:
        branding.esi_gross_threshold = max(0.0, float(payload.esi_gross_threshold))
    if payload.pt_monthly_default is not None:
        branding.pt_monthly_default = max(0.0, float(payload.pt_monthly_default))
    if payload.enable_pf_ceiling is not None:
        branding.enable_pf_ceiling = bool(payload.enable_pf_ceiling)

    db.commit()
    db.refresh(branding)

    return {
        "status": "success",
        "message": "Indian statutory rules and payroll settings updated successfully.",
        "branding": branding.to_dict(),
        "config": {
            "payroll_structure": branding.payroll_structure or "STRUCTURED_SALARY",
            "default_hourly_rate": float(branding.default_hourly_rate or 15.0),
            "standard_working_hours_per_day": float(branding.standard_working_hours_per_day or 8.0),
            "enable_overtime": bool(branding.enable_overtime),
            "overtime_rate_multiplier": float(branding.overtime_rate_multiplier or 1.5),
            "holiday_ot_multiplier": float(branding.holiday_ot_multiplier or 2.0),
            "missed_checkout_policy": branding.missed_checkout_policy or "HALF_DAY",
            "currency_symbol": branding.currency_symbol or "₹",
            "epf_employee_pct": float(branding.epf_employee_pct or 12.0),
            "epf_employer_pct": float(branding.epf_employer_pct or 12.0),
            "epf_wage_ceiling": float(branding.epf_ceiling_limit or 15000.0),
            "epf_admin_charges_pct": float(branding.epf_admin_charges_pct or 0.50),
            "esi_employee_pct": float(branding.esic_employee_pct or 0.75),
            "esi_employer_pct": float(branding.esic_employer_pct or 3.25),
            "esi_gross_threshold": float(branding.esi_gross_threshold or 21000.0),
            "pt_monthly_default": float(branding.pt_monthly_default or 200.0),
            "enable_pf_ceiling": bool(branding.enable_pf_ceiling),
        },
    }


# ==============================================================================
# 6. SUMMARY & EXPORT
# ==============================================================================

@router.get("/summary")
def get_payroll_summary_api(
    start_date: Optional[str] = Query(None, description="Start date (YYYY-MM-DD)"),
    end_date: Optional[str] = Query(None, description="End date (YYYY-MM-DD)"),
    department_id: Optional[int] = Query(None),
    user_role: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    now = get_ist_now()
    if start_date:
        try:
            s_date = datetime.strptime(start_date, "%Y-%m-%d").date()
        except ValueError:
            s_date = now.date() - timedelta(days=7)
    else:
        s_date = now.date() - timedelta(days=7)

    if end_date:
        try:
            e_date = datetime.strptime(end_date, "%Y-%m-%d").date()
        except ValueError:
            e_date = now.date()
    else:
        e_date = now.date()

    payroll_service = PayrollService(db)
    return payroll_service.get_payroll_summary(
        tenant=current_tenant,
        start_date=s_date,
        end_date=e_date,
        department_id=department_id,
        user_role=user_role,
        search_term=search,
    )


@router.get("/export")
def export_payroll_api(
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    export_format: str = Query("xlsx", pattern="^(csv|xlsx)$"),
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    now = get_ist_now()
    if start_date:
        try:
            s_date = datetime.strptime(start_date, "%Y-%m-%d").date()
        except ValueError:
            s_date = now.date() - timedelta(days=30)
    else:
        s_date = now.date() - timedelta(days=30)

    if end_date:
        try:
            e_date = datetime.strptime(end_date, "%Y-%m-%d").date()
        except ValueError:
            e_date = now.date()
    else:
        e_date = now.date()

    payroll_service = PayrollService(db)
    return payroll_service.export_payroll(
        tenant=current_tenant,
        start_date=s_date,
        end_date=e_date,
        export_format=export_format,
    )


@router.post("/employee-rate")
def update_employee_rate_api(
    payload: EmployeeRateUpdateRequest,
    db: Session = Depends(get_db),
    current_tenant: Tenant = Depends(get_current_payroll_tenant),
):
    check_tenant_operational_access(current_tenant)
    payroll_service = PayrollService(db)
    return payroll_service.update_employee_rate(
        tenant_id=current_tenant.id,
        student_id=payload.student_id,
        hourly_rate=payload.hourly_rate,
        monthly_base_salary=payload.monthly_base_salary,
    )


def calculate_payroll_data(
    db: Session,
    tenant: Tenant,
    start_date: date,
    end_date: date,
    department_id: Optional[int] = None,
    user_role: Optional[str] = None,
    search_term: Optional[str] = None,
) -> Dict[str, Any]:
    """Helper function maintaining backward compatibility with legacy tests/callers."""
    payroll_service = PayrollService(db)
    return payroll_service.get_payroll_summary(
        tenant=tenant,
        start_date=start_date,
        end_date=end_date,
        department_id=department_id,
        user_role=user_role,
        search_term=search_term,
    )

