"""
Payroll Service
Handles enterprise payroll batch calculation, salary structure assignment & revisions,
statutory deductions (EPF, ESIC, PT, TDS), payslips, bank advice, and export generation.
"""

import io
import json
import logging
from datetime import datetime, date, timedelta
from typing import Optional, List, Dict, Any, Tuple
import pandas as pd
from sqlalchemy.orm import Session, joinedload
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from fastapi import HTTPException, status
from fastapi.responses import StreamingResponse

from src.database.models import (
    Tenant,
    Student,
    AttendanceRecord,
    SystemBranding,
    Department,
    CompanyLocation,
    DesignationMaster,
    SalaryTemplate,
    EmployeeSalaryStructure,
    SalaryRevisionHistory,
    PayrollBatch,
    PayrollPayslip,
    User,
)
from src.core.payroll_engine import (
    calculate_employee_payroll,
    number_to_words_inr,
    get_effective_salary_structures,
)
from src.utils.timezone import get_ist_now, get_ist_date

logger = logging.getLogger("payroll_service")


class PayrollService:
    def __init__(self, db: Session):
        self.db = db

    # =========================================================================
    # 1. EMPLOYEE STRUCTURES & OVERVIEW
    # =========================================================================
    def get_structure_overview(self, tenant: Tenant, department_id: Optional[int] = None) -> List[Dict[str, Any]]:
        query = (
            self.db.query(Student)
            .filter(Student.tenant_id == tenant.id, Student.is_active == True)
            .options(
                joinedload(Student.department_rel),
                joinedload(Student.designation_rel),
                joinedload(Student.location),
            )
        )
        if department_id:
            query = query.filter(Student.department_id == department_id)

        employees = query.order_by(Student.name.asc()).all()
        result = []
        for emp in employees:
            current_structure = (
                self.db.query(EmployeeSalaryStructure)
                .filter(
                    EmployeeSalaryStructure.tenant_id == tenant.id,
                    EmployeeSalaryStructure.student_id == emp.id,
                    EmployeeSalaryStructure.is_current == True,
                )
                .first()
            )
            result.append({
                "id": emp.id,
                "roll_number": emp.roll_number,
                "name": emp.name,
                "department": emp.department or (emp.department_rel.name if emp.department_rel else "General"),
                "department_id": emp.department_id,
                "designation": emp.designation or (emp.designation_rel.title if emp.designation_rel else "Staff"),
                "designation_id": emp.designation_id,
                "default_template_id": emp.designation_rel.salary_template_id if emp.designation_rel else None,
                "location_name": emp.location.name if emp.location else "Head Office",
                "location_id": emp.location_id,
                "pan_number": emp.pan_number or "",
                "uan_number": emp.uan_number or "",
                "esic_number": emp.esic_number or "",
                "bank_name": emp.bank_name or "",
                "bank_account_number": emp.bank_account_number or "",
                "bank_ifsc_code": emp.bank_ifsc_code or "",
                "hourly_rate": float(emp.hourly_rate or 0.0),
                "monthly_base_salary": float(emp.monthly_base_salary or 0.0),
                "active_structure": current_structure.to_dict() if current_structure else None,
            })
        return result

    def get_employee_structure(self, tenant_id: int, student_id: int) -> Dict[str, Any]:
        emp = (
            self.db.query(Student)
            .filter(Student.id == student_id, Student.tenant_id == tenant_id)
            .options(
                joinedload(Student.department_rel),
                joinedload(Student.designation_rel),
                joinedload(Student.location),
            )
            .first()
        )
        if not emp:
            raise HTTPException(status_code=404, detail="Employee not found.")

        current_structure = (
            self.db.query(EmployeeSalaryStructure)
            .filter(
                EmployeeSalaryStructure.tenant_id == tenant_id,
                EmployeeSalaryStructure.student_id == student_id,
                EmployeeSalaryStructure.is_current == True,
            )
            .first()
        )

        history = (
            self.db.query(SalaryRevisionHistory)
            .filter(
                SalaryRevisionHistory.tenant_id == tenant_id,
                SalaryRevisionHistory.student_id == student_id,
            )
            .order_by(SalaryRevisionHistory.effective_from_date.desc(), SalaryRevisionHistory.created_at.desc())
            .all()
        )

        return {
            "status": "success",
            "employee": {
                "id": emp.id,
                "roll_number": emp.roll_number,
                "name": emp.name,
                "department": emp.department or (emp.department_rel.name if emp.department_rel else "General"),
                "department_id": emp.department_id,
                "designation": emp.designation or (emp.designation_rel.title if emp.designation_rel else "Staff"),
                "designation_id": emp.designation_id,
                "default_template_id": emp.designation_rel.salary_template_id if emp.designation_rel else None,
                "location_name": emp.location.name if emp.location else "Head Office",
                "location_id": emp.location_id,
                "pan_number": emp.pan_number or "",
                "uan_number": emp.uan_number or "",
                "esic_number": emp.esic_number or "",
                "bank_name": emp.bank_name or "",
                "bank_account_number": emp.bank_account_number or "",
                "bank_ifsc_code": emp.bank_ifsc_code or "",
                "monthly_base_salary": float(emp.monthly_base_salary or 0.0),
                "hourly_rate": float(emp.hourly_rate or 0.0),
            },
            "current_structure": current_structure.to_dict() if current_structure else None,
            "revision_history": [h.to_dict() for h in history],
        }

    def assign_salary_structure(self, tenant_id: int, payload: Any, current_user: Optional[User] = None) -> Dict[str, Any]:
        emp = (
            self.db.query(Student)
            .filter(Student.id == payload.student_id, Student.tenant_id == tenant_id)
            .first()
        )
        if not emp:
            raise HTTPException(status_code=404, detail="Employee not found.")

        try:
            eff_from = datetime.strptime(payload.effective_from_date, "%Y-%m-%d").date()
        except ValueError:
            eff_from = get_ist_date()

        model = payload.compensation_model.upper()
        basic_pct = 50.0
        hra_pct = 20.0
        da_pct = 0.0
        conv_val = 1600.0
        med_val = 1250.0
        enable_pf = payload.enable_pf if payload.enable_pf is not None else True
        pf_capped = payload.pf_capped_at_ceiling if payload.pf_capped_at_ceiling is not None else True
        enable_esi = payload.enable_esi if payload.enable_esi is not None else True
        enable_pt = payload.enable_pt if payload.enable_pt is not None else True

        if payload.template_id:
            tpl = (
                self.db.query(SalaryTemplate)
                .filter(SalaryTemplate.id == payload.template_id, SalaryTemplate.tenant_id == tenant_id)
                .first()
            )
            if tpl:
                model = tpl.compensation_model
                basic_pct = tpl.basic_percentage
                hra_pct = tpl.hra_percentage
                da_pct = tpl.da_percentage
                conv_val = tpl.conveyance_fixed
                med_val = tpl.medical_fixed
                enable_pf = tpl.enable_pf
                pf_capped = tpl.pf_capped_at_ceiling
                enable_esi = tpl.enable_esi
                enable_pt = tpl.enable_pt

        annual_ctc = float(payload.annual_ctc or 0.0)
        monthly_gross = float(payload.monthly_gross or 0.0)

        if annual_ctc > 0 and monthly_gross <= 0:
            monthly_gross = round(annual_ctc / 12.0, 2)
        elif monthly_gross > 0 and annual_ctc <= 0:
            annual_ctc = round(monthly_gross * 12.0, 2)

        if model == "STRUCTURED_SALARY" and monthly_gross > 0:
            monthly_basic = float(payload.monthly_basic or round(monthly_gross * (basic_pct / 100.0), 2))
            monthly_da = float(payload.monthly_da or round(monthly_gross * (da_pct / 100.0), 2))
            monthly_hra = float(payload.monthly_hra or round(monthly_basic * (hra_pct / 100.0), 2))
            conv_allow = float(payload.conveyance_allowance or conv_val)
            med_allow = float(payload.medical_allowance or med_val)
            specified_total = monthly_basic + monthly_da + monthly_hra + conv_allow + med_allow
            special_allow = float(payload.special_allowance or max(0.0, round(monthly_gross - specified_total, 2)))
            other_allow = float(payload.other_allowances or 0.0)
        else:
            monthly_basic = float(payload.monthly_basic or (monthly_gross * 0.50 if monthly_gross > 0 else 0.0))
            monthly_da = float(payload.monthly_da or 0.0)
            monthly_hra = float(payload.monthly_hra or 0.0)
            conv_allow = float(payload.conveyance_allowance or 0.0)
            med_allow = float(payload.medical_allowance or 0.0)
            special_allow = float(payload.special_allowance or 0.0)
            other_allow = float(payload.other_allowances or 0.0)

        existing_current = (
            self.db.query(EmployeeSalaryStructure)
            .filter(
                EmployeeSalaryStructure.tenant_id == tenant_id,
                EmployeeSalaryStructure.student_id == emp.id,
                EmployeeSalaryStructure.is_current == True,
            )
            .first()
        )

        prev_ctc = existing_current.annual_ctc if existing_current else float(emp.monthly_base_salary or 0.0) * 12.0
        prev_gross = existing_current.monthly_gross if existing_current else float(emp.monthly_base_salary or 0.0)
        prev_model = existing_current.compensation_model if existing_current else "NONE"

        if existing_current:
            existing_current.is_current = False
            if not existing_current.effective_to_date or existing_current.effective_to_date >= eff_from:
                existing_current.effective_to_date = eff_from - timedelta(days=1)

        new_structure = EmployeeSalaryStructure(
            tenant_id=tenant_id,
            student_id=emp.id,
            template_id=payload.template_id,
            compensation_model=model,
            annual_ctc=annual_ctc,
            monthly_gross=monthly_gross,
            monthly_basic=monthly_basic,
            monthly_da=monthly_da,
            monthly_hra=monthly_hra,
            conveyance_allowance=conv_allow,
            medical_allowance=med_allow,
            special_allowance=special_allow,
            other_allowances=other_allow,
            hourly_rate=float(payload.hourly_rate or emp.hourly_rate or 0.0),
            daily_rate=float(payload.daily_rate or (round(monthly_gross / 26.0, 2) if monthly_gross > 0 else 0.0)),
            fixed_stipend=float(payload.fixed_stipend or (monthly_gross if model == "STIPEND" else 0.0)),
            commission_percentage=float(payload.commission_percentage or 0.0),
            enable_pf=enable_pf,
            pf_capped_at_ceiling=pf_capped,
            enable_esi=enable_esi,
            enable_pt=enable_pt,
            pt_monthly_amount=float(payload.pt_monthly_amount or 200.0),
            tds_monthly_amount=float(payload.tds_monthly_amount or 0.0),
            effective_from_date=eff_from,
            effective_to_date=None,
            is_current=True,
            revision_reason=payload.revision_reason or "Salary Revision",
            revised_by_user_id=current_user.id if current_user else None,
        )
        self.db.add(new_structure)
        self.db.flush()

        if monthly_gross > 0:
            emp.monthly_base_salary = monthly_gross
        if float(payload.hourly_rate or 0.0) > 0:
            emp.hourly_rate = float(payload.hourly_rate)

        rev_history = SalaryRevisionHistory(
            tenant_id=tenant_id,
            student_id=emp.id,
            salary_structure_id=new_structure.id,
            effective_from_date=eff_from,
            effective_to_date=None,
            previous_annual_ctc=prev_ctc,
            new_annual_ctc=annual_ctc,
            previous_monthly_gross=prev_gross,
            new_monthly_gross=monthly_gross,
            previous_model=prev_model,
            new_model=model,
            revision_reason=payload.revision_reason or "Salary Revision",
            revised_by_user_id=current_user.id if current_user else None,
        )
        self.db.add(rev_history)
        self.db.commit()
        self.db.refresh(new_structure)

        return {
            "status": "success",
            "message": f"Salary structure assigned successfully to {emp.name} ({emp.roll_number}).",
            "data": new_structure.to_dict(),
        }

    def update_statutory_and_banking(self, tenant_id: int, student_id: int, payload: Any) -> Dict[str, Any]:
        emp = (
            self.db.query(Student)
            .filter(Student.id == student_id, Student.tenant_id == tenant_id)
            .first()
        )
        if not emp:
            raise HTTPException(status_code=404, detail="Employee not found.")

        if payload.pan_number is not None:
            emp.pan_number = payload.pan_number.strip().upper() if payload.pan_number else None
        if payload.uan_number is not None:
            emp.uan_number = payload.uan_number.strip() if payload.uan_number else None
        if payload.esic_number is not None:
            emp.esic_number = payload.esic_number.strip() if payload.esic_number else None
        if payload.bank_name is not None:
            emp.bank_name = payload.bank_name.strip() if payload.bank_name else None
        if payload.bank_account_number is not None:
            emp.bank_account_number = payload.bank_account_number.strip() if payload.bank_account_number else None
        if payload.bank_ifsc_code is not None:
            emp.bank_ifsc_code = payload.bank_ifsc_code.strip().upper() if payload.bank_ifsc_code else None
        if payload.location_id is not None:
            emp.location_id = payload.location_id
        if payload.designation_id is not None:
            emp.designation_id = payload.designation_id
            desig = self.db.query(DesignationMaster).filter(DesignationMaster.id == payload.designation_id).first()
            if desig:
                emp.designation = desig.title

        self.db.commit()
        self.db.refresh(emp)

        return {
            "status": "success",
            "message": f"Statutory and banking details updated for {emp.name}.",
            "data": {
                "pan_number": emp.pan_number,
                "uan_number": emp.uan_number,
                "esic_number": emp.esic_number,
                "bank_name": emp.bank_name,
                "bank_account_number": emp.bank_account_number,
                "bank_ifsc_code": emp.bank_ifsc_code,
                "location_id": emp.location_id,
                "designation_id": emp.designation_id,
            },
        }

    # =========================================================================
    # 2. BATCHES & MONTHLY RUN LIFECYCLE
    # =========================================================================
    def list_batches(self, tenant_id: int, year: Optional[int] = None, status_filter: Optional[str] = None) -> List[Dict[str, Any]]:
        query = self.db.query(PayrollBatch).filter(PayrollBatch.tenant_id == tenant_id)
        if year:
            query = query.filter(PayrollBatch.period_year == year)
        if status_filter:
            query = query.filter(PayrollBatch.status == status_filter.strip().upper())

        batches = query.order_by(PayrollBatch.period_year.desc(), PayrollBatch.period_month.desc()).all()
        return [b.to_dict() for b in batches]

    def generate_batch(self, tenant: Tenant, payload: Any, current_user: Optional[User] = None) -> Dict[str, Any]:
        p_year = payload.period_year
        p_month = payload.period_month

        if payload.start_date and payload.end_date:
            try:
                start_date = datetime.strptime(payload.start_date, "%Y-%m-%d").date()
                end_date = datetime.strptime(payload.end_date, "%Y-%m-%d").date()
            except ValueError:
                raise HTTPException(status_code=400, detail="Invalid date format. Expected YYYY-MM-DD.")
        else:
            start_date = date(p_year, p_month, 1)
            if p_month == 12:
                end_date = date(p_year, 12, 31)
            else:
                next_month_start = date(p_year, p_month + 1, 1)
                end_date = next_month_start - timedelta(days=1)

        working_days = float(payload.working_days or 26.0)
        month_code = f"{p_year}-{p_month:02d}"

        existing_batch = (
            self.db.query(PayrollBatch)
            .filter(
                PayrollBatch.tenant_id == tenant.id,
                PayrollBatch.period_year == p_year,
                PayrollBatch.period_month == p_month,
            )
            .first()
        )

        if existing_batch and existing_batch.status in ("APPROVED", "DISBURSED", "LOCKED"):
            raise HTTPException(
                status_code=400,
                detail=f"Payroll batch for {month_code} is already {existing_batch.status} and locked for re-generation.",
            )

        if existing_batch:
            self.db.delete(existing_batch)
            self.db.flush()

        emp_query = (
            self.db.query(Student)
            .filter(Student.tenant_id == tenant.id, Student.is_active == True)
        )
        if payload.department_id:
            emp_query = emp_query.filter(Student.department_id == payload.department_id)

        employees = emp_query.order_by(Student.name.asc()).all()
        if not employees:
            raise HTTPException(status_code=400, detail="No active employees found to generate payroll.")

        batch_number = f"PAY-{tenant.slug.upper()[:4]}-{p_year}{p_month:02d}"
        batch = PayrollBatch(
            tenant_id=tenant.id,
            batch_number=batch_number,
            period_month=p_month,
            period_year=p_year,
            start_date=start_date,
            end_date=end_date,
            total_working_days=working_days,
            status="DRAFT",
            total_employees_count=len(employees),
            processed_by_user_id=current_user.id if current_user else None,
            processed_at=get_ist_now(),
        )
        self.db.add(batch)
        self.db.flush()

        batch_gross = 0.0
        batch_net = 0.0
        batch_pf = 0.0
        batch_esi = 0.0
        batch_pt = 0.0
        batch_tds = 0.0
        batch_employer_liability = 0.0

        for emp in employees:
            res = calculate_employee_payroll(
                db=self.db,
                tenant=tenant,
                student=emp,
                start_date=start_date,
                end_date=end_date,
                total_working_days=working_days,
            )

            payslip = PayrollPayslip(
                tenant_id=tenant.id,
                batch_id=batch.id,
                student_id=emp.id,
                template_id=res.get("template_id"),
                period_month=p_month,
                period_year=p_year,
                calendar_days=res.get("calendar_days", (end_date - start_date).days + 1),
                working_days=res.get("working_days", working_days),
                present_days=res.get("present_days", 0.0),
                paid_leave_days=res.get("paid_leave_days", 0.0),
                unpaid_leave_days=res.get("unpaid_leave_days", 0.0),
                absent_days=res.get("absent_days", 0.0),
                billable_hours=res.get("billable_hours", 0.0),
                regular_ot_hours=res.get("regular_ot_hours", 0.0),
                holiday_ot_hours=res.get("holiday_ot_hours", 0.0),
                ot_earnings=res.get("ot_earnings", 0.0),
                basic_earned=res.get("basic_earned", 0.0),
                da_earned=res.get("da_earned", 0.0),
                hra_earned=res.get("hra_earned", 0.0),
                conveyance_earned=res.get("conveyance_earned", 0.0),
                medical_earned=res.get("medical_earned", 0.0),
                special_allowance_earned=res.get("special_allowance_earned", 0.0),
                other_earnings=res.get("other_earnings", 0.0),
                incentives_bonus=res.get("incentives_bonus", 0.0),
                gross_earnings=res.get("gross_earnings", 0.0),
                epf_employee=res.get("epf_employee", 0.0),
                epf_employer=res.get("epf_employer", 0.0),
                eps_employer=res.get("eps_employer", 0.0),
                esic_employee=res.get("esic_employee", 0.0),
                esic_employer=res.get("esic_employer", 0.0),
                professional_tax=res.get("professional_tax", 0.0),
                tds_deduction=res.get("tds_deduction", 0.0),
                advance_loan_deduction=res.get("advance_loan_deduction", 0.0),
                other_deductions=res.get("other_deductions", 0.0),
                total_deductions=res.get("total_deductions", 0.0),
                net_salary=res.get("net_salary", 0.0),
                employer_total_ctc_outlay=res.get("employer_total_ctc_outlay", 0.0),
                payment_status="PENDING",
                payment_method=res.get("payment_method", "BANK_TRANSFER"),
                breakdown_json=json.dumps(res.get("breakdown", {})),
                remarks=res.get("remarks", ""),
            )
            self.db.add(payslip)

            batch_gross += res.get("gross_earnings", 0.0)
            batch_net += res.get("net_salary", 0.0)
            batch_pf += res.get("epf_employee", 0.0) + res.get("epf_employer", 0.0) + res.get("eps_employer", 0.0)
            batch_esi += res.get("esic_employee", 0.0) + res.get("esic_employer", 0.0)
            batch_pt += res.get("professional_tax", 0.0)
            batch_tds += res.get("tds_deduction", 0.0)
            batch_employer_liability += (
                res.get("epf_employer", 0.0) + res.get("eps_employer", 0.0) + res.get("esic_employer", 0.0)
            )

        batch.total_gross_outlay = round(batch_gross, 2)
        batch.total_net_outlay = round(batch_net, 2)
        batch.total_pf_liability = round(batch_pf, 2)
        batch.total_esi_liability = round(batch_esi, 2)
        batch.total_pt_liability = round(batch_pt, 2)
        batch.total_tds_liability = round(batch_tds, 2)
        batch.total_employer_contributions = round(batch_employer_liability, 2)

        self.db.commit()
        self.db.refresh(batch)

        return {
            "status": "success",
            "message": f"Payroll batch '{batch.batch_number}' generated for {len(employees)} employees.",
            "data": batch.to_dict(),
        }

    def get_batch_details(self, tenant_id: int, batch_id: int) -> Dict[str, Any]:
        batch = (
            self.db.query(PayrollBatch)
            .filter(PayrollBatch.id == batch_id, PayrollBatch.tenant_id == tenant_id)
            .first()
        )
        if not batch:
            raise HTTPException(status_code=404, detail="Payroll batch not found.")

        payslips = (
            self.db.query(PayrollPayslip)
            .filter(PayrollPayslip.batch_id == batch.id, PayrollPayslip.tenant_id == tenant_id)
            .order_by(PayrollPayslip.id.asc())
            .all()
        )

        return {
            "status": "success",
            "batch": batch.to_dict(),
            "payslips": [p.to_dict() for p in payslips],
        }

    def update_batch_status(self, tenant_id: int, batch_id: int, action: str, current_user: Optional[User] = None) -> Dict[str, Any]:
        batch = (
            self.db.query(PayrollBatch)
            .filter(PayrollBatch.id == batch_id, PayrollBatch.tenant_id == tenant_id)
            .first()
        )
        if not batch:
            raise HTTPException(status_code=404, detail="Payroll batch not found.")

        if action == "VERIFY":
            batch.status = "VERIFIED"
        elif action == "APPROVE":
            batch.status = "APPROVED"
            batch.approved_by_user_id = current_user.id if current_user else None
            batch.approved_at = get_ist_now()
        elif action == "DISBURSE":
            batch.status = "DISBURSED"
            batch.disbursed_by_user_id = current_user.id if current_user else None
            batch.disbursed_at = get_ist_now()

            payslips = (
                self.db.query(PayrollPayslip)
                .filter(PayrollPayslip.batch_id == batch.id, PayrollPayslip.tenant_id == tenant_id)
                .all()
            )
            for p in payslips:
                p.payment_status = "PAID"
                p.payment_date = get_ist_date()

        self.db.commit()
        self.db.refresh(batch)
        return {
            "status": "success",
            "message": f"Batch {batch.batch_number} marked as {batch.status}.",
            "data": batch.to_dict(),
        }

    def delete_batch(self, tenant_id: int, batch_id: int) -> Dict[str, Any]:
        batch = (
            self.db.query(PayrollBatch)
            .filter(PayrollBatch.id == batch_id, PayrollBatch.tenant_id == tenant_id)
            .first()
        )
        if not batch:
            raise HTTPException(status_code=404, detail="Payroll batch not found.")

        if batch.status in ("APPROVED", "DISBURSED", "LOCKED"):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Cannot delete batch '{batch.batch_number}' because it is in '{batch.status}' status.",
            )

        self.db.delete(batch)
        self.db.commit()
        return {"status": "success", "message": f"Payroll batch '{batch.batch_number}' deleted successfully."}

    # =========================================================================
    # 3. INDIVIDUAL PAYSLIP & SUMMARY
    # =========================================================================
    def get_payslip(self, tenant: Tenant, payslip_id: int) -> Dict[str, Any]:
        payslip = (
            self.db.query(PayrollPayslip)
            .filter(PayrollPayslip.id == payslip_id, PayrollPayslip.tenant_id == tenant.id)
            .first()
        )
        if not payslip:
            raise HTTPException(status_code=404, detail="Payslip not found.")

        data = payslip.to_dict()
        data["net_in_words"] = number_to_words_inr(payslip.net_salary)
        branding_dict = tenant.branding.to_dict() if tenant.branding else {}
        data["company_name"] = branding_dict.get("institution_name") or tenant.name
        data["company_logo"] = branding_dict.get("logo_url") or ""
        data["currency_symbol"] = branding_dict.get("currency_symbol") or "₹"
        return {"status": "success", "data": data}

    # =========================================================================
    # 4. HOURLY / SUMMARY & EXPORT
    # =========================================================================
    def get_payroll_summary(
        self,
        tenant: Tenant,
        start_date: date,
        end_date: date,
        department_id: Optional[int] = None,
        user_role: Optional[str] = None,
        search_term: Optional[str] = None,
    ) -> Dict[str, Any]:
        branding = tenant.branding
        currency = branding.currency_symbol if branding and branding.currency_symbol else "₹"
        working_days = float((end_date - start_date).days + 1)
        std_daily_hours = float(branding.standard_working_hours_per_day or 8.0) if branding else 8.0
        enable_ot = bool(branding.enable_overtime) if branding else True
        ot_mult = float(branding.overtime_rate_multiplier or 1.5) if branding else 1.5
        missed_policy = (branding.missed_checkout_policy or "HALF_DAY") if branding else "HALF_DAY"

        emp_query = (
            self.db.query(Student)
            .filter(Student.tenant_id == tenant.id, Student.is_active == True)
            .options(joinedload(Student.department_rel), joinedload(Student.designation_rel))
        )
        if department_id:
            emp_query = emp_query.filter(Student.department_id == department_id)
        if user_role:
            emp_query = emp_query.filter(Student.user_role == user_role.strip().lower())
        if search_term:
            term = f"%{search_term.strip()}%"
            emp_query = emp_query.filter((Student.name.ilike(term)) | (Student.roll_number.ilike(term)))

        employees = emp_query.order_by(Student.name.asc()).all()

        start_dt = datetime.combine(start_date, datetime.min.time())
        end_dt = datetime.combine(end_date, datetime.max.time())

        items = []
        total_billed_hours = 0.0
        total_ot_hours = 0.0
        total_gross_wages = 0.0
        total_estimated_payout = 0.0

        for emp in employees:
            rate = float(emp.hourly_rate or (branding.default_hourly_rate if branding else 15.0) or 15.0)

            # Query attendance records in the window
            recs = (
                self.db.query(AttendanceRecord)
                .filter(
                    AttendanceRecord.tenant_id == tenant.id,
                    AttendanceRecord.student_id == emp.id,
                    AttendanceRecord.timestamp.between(start_dt, end_dt),
                )
                .all()
            )

            emp_shift = emp.shift
            emp_std_hours = float(emp_shift.total_shift_hours if emp_shift else std_daily_hours)

            daily_hours_map: Dict[date, float] = {}
            missed_count = 0
            for r in recs:
                r_date = r.timestamp.date() if r.timestamp else start_date
                if r.check_out_time and r.work_duration_minutes is not None:
                    h = max(0.0, r.work_duration_minutes / 60.0)
                elif r.check_in_time and not r.check_out_time:
                    missed_count += 1
                    if missed_policy == "HALF_DAY":
                        h = emp_std_hours / 2.0
                    elif missed_policy == "ZERO":
                        h = 0.0
                    else:
                        h = emp_std_hours
                else:
                    h = 0.0
                daily_hours_map[r_date] = daily_hours_map.get(r_date, 0.0) + h

            emp_std_h = 0.0
            emp_ot_h = 0.0
            for d_date, d_h in daily_hours_map.items():
                if enable_ot:
                    std_part = min(d_h, emp_std_hours)
                    ot_part = max(0.0, d_h - emp_std_hours)
                else:
                    std_part = d_h
                    ot_part = 0.0
                emp_std_h += std_part
                emp_ot_h += ot_part

            res = calculate_employee_payroll(
                db=self.db,
                tenant=tenant,
                student=emp,
                start_date=start_date,
                end_date=end_date,
                total_working_days=working_days,
            )

            paid_leave_days = float(res.get("paid_leave_days", 0.0))
            paid_leave_hours = round(paid_leave_days * emp_std_hours, 2)

            emp_std_h = emp_std_h + paid_leave_hours
            total_h = round(emp_std_h + emp_ot_h, 2)
            std_pay = round(emp_std_h * rate, 2)
            ot_pay = round(emp_ot_h * rate * ot_mult, 2) if enable_ot else 0.0
            gross_pay = round(std_pay + ot_pay, 2)

            if res.get("compensation_model") == "STRUCTURED_SALARY" and float(emp.monthly_base_salary or 0.0) > 0:
                final_gross = res.get("gross_earnings", gross_pay)
                final_net = res.get("net_salary", gross_pay)
            else:
                final_gross = gross_pay
                final_net = gross_pay

            total_billed_hours += total_h
            total_ot_hours += emp_ot_h
            total_gross_wages += final_gross
            total_estimated_payout += final_net

            items.append({
                "student_id": emp.id,
                "roll_number": emp.roll_number,
                "name": emp.name,
                "department": emp.department or (emp.department_rel.name if emp.department_rel else "General"),
                "designation": emp.designation or (emp.designation_rel.title if emp.designation_rel else "Staff"),
                "user_role": emp.user_role,
                "hourly_rate": rate,
                "monthly_base_salary": float(emp.monthly_base_salary or 0.0),
                "total_hours": total_h,
                "total_active_hours": total_h,
                "standard_hours": emp_std_h,
                "regular_hours": emp_std_h,
                "overtime_hours": emp_ot_h,
                "paid_leave_days": paid_leave_days,
                "paid_leave_hours": paid_leave_hours,
                "standard_pay": std_pay,
                "overtime_pay": ot_pay,
                "total_days_present": res.get("present_days", 0.0),
                "days_worked": len([d for d, h in daily_hours_map.items() if h > 0]) or res.get("present_days", 0.0),
                "gross_pay": final_gross,
                "net_pay": final_net,
                "missed_checkout_count": missed_count,
                "epf_deduction": res.get("epf_employee", 0.0),
                "esic_deduction": res.get("esic_employee", 0.0),
                "pt_deduction": res.get("professional_tax", 0.0),
                "tds_deduction": res.get("tds_deduction", 0.0),
                "model": res.get("compensation_model", "STRUCTURED_SALARY"),
            })

        return {
            "status": "success",
            "date_range": {"start_date": str(start_date), "end_date": str(end_date)},
            "summary": {
                "total_employees": len(employees),
                "total_hours": round(total_billed_hours, 2),
                "total_overtime_hours": round(total_ot_hours, 2),
                "total_gross_wages": round(total_gross_wages, 2),
                "total_estimated_payout": round(total_estimated_payout, 2),
                "currency_symbol": currency,
            },
            "items": items,
            "records": items,
        }

    def export_payroll(self, tenant: Tenant, start_date: date, end_date: date, export_format: str = "csv") -> StreamingResponse:
        data = self.get_payroll_summary(tenant, start_date, end_date)
        records = data.get("records", [])
        branding = tenant.branding
        currency = branding.currency_symbol if branding and branding.currency_symbol else "₹"
        company_name = branding.institution_name if branding else tenant.name

        rows = []
        for r in records:
            rows.append({
                "Emp Code": r["roll_number"],
                "Employee Name": r["name"],
                "Department": r["department"],
                "Designation": r["designation"],
                "Compensation Model": r["model"],
                "Days Present": r["total_days_present"],
                "Billable Hours": r["total_hours"],
                "OT Hours": r["overtime_hours"],
                "Gross Earnings": r["gross_pay"],
                "EPF Employee": r["epf_deduction"],
                "ESIC Employee": r["esic_deduction"],
                "Professional Tax": r["pt_deduction"],
                "TDS": r["tds_deduction"],
                "Net Payout": r["net_pay"],
            })

        df = pd.DataFrame(rows)
        filename = f"payroll_summary_{tenant.slug}_{start_date}_{end_date}"

        if export_format == "csv":
            csv_buf = io.StringIO()
            csv_buf.write(f"# COMPANY: {company_name}\n")
            csv_buf.write(f"# PERIOD: {start_date} to {end_date}\n")
            csv_buf.write(f"# TOTAL PAYOUT: {currency}{data['summary']['total_estimated_payout']:.2f}\n")
            df.to_csv(csv_buf, index=False)
            csv_buf.seek(0)
            return StreamingResponse(
                io.BytesIO(csv_buf.getvalue().encode("utf-8")),
                media_type="text/csv",
                headers={"Content-Disposition": f"attachment; filename={filename}.csv"},
            )
        else:
            excel_buf = io.BytesIO()
            with pd.ExcelWriter(excel_buf, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Payroll Summary", startrow=3)
                ws = writer.sheets["Payroll Summary"]
                ws.merge_cells("A1:N1")
                ws["A1"] = f"{company_name} — Payroll & Statutory Summary ({start_date} to {end_date})"
                ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
                ws["A1"].fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
                ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
            excel_buf.seek(0)
            return StreamingResponse(
                excel_buf,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"},
            )

    # =========================================================================
    # 5. LEGACY / EMPLOYEE RATE & SETTINGS
    # =========================================================================
    def update_employee_rate(self, tenant_id: int, student_id: int, hourly_rate: Optional[float], monthly_base_salary: Optional[float]) -> Dict[str, Any]:
        student = (
            self.db.query(Student)
            .filter(Student.id == student_id, Student.tenant_id == tenant_id)
            .first()
        )
        if not student:
            raise HTTPException(status_code=404, detail="Employee not found.")

        if hourly_rate is not None:
            student.hourly_rate = max(0.0, float(hourly_rate))
        if monthly_base_salary is not None:
            student.monthly_base_salary = max(0.0, float(monthly_base_salary))

        self.db.commit()
        self.db.refresh(student)

        return {
            "status": "success",
            "message": f"Updated wage rates for {student.name}.",
            "data": {
                "student_id": student.id,
                "hourly_rate": float(student.hourly_rate or 0.0),
                "monthly_base_salary": float(student.monthly_base_salary or 0.0),
            },
        }
